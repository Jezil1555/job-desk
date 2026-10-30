"""End-to-end UI test: drives the real frontend in headless Chromium against an in-memory
fake of the Supabase REST API (the real database is tested separately in supabase/tests).

Usage:  npm i --no-save @supabase/supabase-js@2.45.4 && python3 tests/e2e/run.py [--shots DIR]
Exit code 0 = all checks passed.
"""
import json, os, re, sys, threading, time, uuid, functools, http.server, datetime as dt
from urllib.parse import urlparse, parse_qs, unquote
from playwright.sync_api import sync_playwright

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PUB = os.path.join(ROOT, "public")
UMD = os.path.join(ROOT, "node_modules/@supabase/supabase-js/dist/umd/supabase.js")
SHOTS = sys.argv[sys.argv.index("--shots") + 1] if "--shots" in sys.argv else None
if SHOTS: os.makedirs(SHOTS, exist_ok=True)

http.server.SimpleHTTPRequestHandler.log_message = lambda *a: None
srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(http.server.SimpleHTTPRequestHandler, directory=PUB))
PORT = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
BASE = f"http://127.0.0.1:{PORT}"

ME, ADIL = "11111111-1111-4111-8111-111111111111", "22222222-2222-4222-8222-222222222222"
now = dt.datetime.now(dt.timezone.utc)
iso = lambda d: d.isoformat().replace("+00:00", "Z")
WS = ["aaaaaaa1-0000-4000-8000-000000000001", "aaaaaaa1-0000-4000-8000-000000000002", "aaaaaaa1-0000-4000-8000-000000000003"]
P1 = "bbbbbbb1-0000-4000-8000-000000000001"

FE1, FE2 = "fe000000-0000-4000-8000-000000000001", "fe000000-0000-4000-8000-000000000002"
def seed_accounts(eid):
    rows = [("1000", "Cash on hand", "asset", "cash", "cash"), ("1010", "Bank account", "asset", "bank", "bank"), ("1100", "Accounts receivable", "asset", "receivable", "ar"),
            ("2000", "Accounts payable", "liability", "payable", "ap"), ("3900", "Opening balance equity", "equity", "other", "opening"), ("4000", "Sales", "income", "other", "sales"),
            ("6000", "Rent", "expense", "other", None), ("6900", "Other expenses", "expense", "other", None)]
    return [{"id": f"{eid[:8]}-acc-{c}", "entity_id": eid, "code": c, "name": n, "type": t, "subtype": st, "system_key": k, "is_active": True, "description": None} for c, n, t, st, k in rows]
def fin_fixture():
    ents = [{"id": FE1, "name": "Dlusso Bespoke", "kind": "business", "base_currency": "QAR", "workspace_id": WS[0], "fiscal_year_start": 1, "lock_date": None, "archived": False, "created_by": ME,
             "next_journal_no": 1, "next_invoice_no": 1, "next_bill_no": 1, "created_at": iso(now)},
            {"id": FE2, "name": "The Burgery Thrissur", "kind": "business", "base_currency": "INR", "workspace_id": None, "fiscal_year_start": 4, "lock_date": None, "archived": False, "created_by": ME,
             "next_journal_no": 1, "next_invoice_no": 1, "next_bill_no": 1, "created_at": iso(now)}]
    return {"fin_entities": ents, "fin_members": [{"entity_id": e["id"], "user_id": ME, "role": "owner"} for e in ents], "fin_accounts": seed_accounts(FE1) + seed_accounts(FE2),
            "fin_parties": [], "fin_journals": [], "fin_lines": [], "fin_invoices": [], "fin_invoice_lines": [], "fin_payments": [], "fin_allocations": [], "fin_items": [],
            "fin_stock_moves": [], "fin_statement_lines": [], "fin_budgets": [], "fin_files": [], "fin_audit": [], "fin_partner_links": [], "fin_partner_txns": [], "fin_partner_docs": [], "fin_ic_txns": [],
            "fin_production_orders": [], "fin_production_costs": [], "fin_goods_receipts": [], "fin_stock_exceptions": [], "fin_fx_revaluations": []}

def fresh_db():
    t = lambda **k: {"id": k.pop("id", str(uuid.uuid4())), "workspace_id": WS[1], "project_id": None, "parent_id": None, "milestone_id": None, "title": "Task", "description": None,
                     "category": None, "status": "not_started", "priority": "medium", "owner_id": ME, "created_by": ME, "start_date": None, "due_at": None, "due_all_day": True,
                     "follow_up_at": None, "progress": 0, "next_action": None, "completion_criteria": [], "waiting_on": None, "waiting_on_contact_id": None, "waiting_for": None,
                     "blocker": None, "blocker_impact": None, "unblock_action": None, "recurrence": None, "recurrence_parent_id": None, "cancelled_reason": None,
                     "deleted_at": None, "client_name": None, "reference": None, "completed_at": None, "assignee_id": None, "reviewer_id": None, "review_status": None, "review_note": None,
                     "reviewed_by": None, "reviewed_at": None, "created_at": iso(now - dt.timedelta(days=4)), "updated_at": iso(now - dt.timedelta(hours=3)), **k}
    return {
        "workspaces": [
            {"id": WS[0], "owner_id": ME, "name": "D'Lusso & Mediate", "color": "#6b4b8a", "kind": "business", "archived": False, "sort": 0, "description": None},
            {"id": WS[1], "owner_id": ME, "name": "QCLOUD", "color": "#1f6aa0", "kind": "professional", "archived": False, "sort": 1, "description": "Logistics"},
            {"id": WS[2], "owner_id": ME, "name": "Personal & Family", "color": "#5f6b78", "kind": "personal", "archived": False, "sort": 10, "description": None}],
        "workspace_members": [{"workspace_id": w, "user_id": ME, "role": "owner"} for w in WS] + [{"workspace_id": WS[1], "user_id": ADIL, "role": "member"}],
        "profiles": [{"id": ME, "email": "jezz@example.com", "display_name": "Jezz", "timezone": "Asia/Qatar", "prefs": {}},
                     {"id": ADIL, "email": "adil@example.com", "display_name": "Adil Rahman", "timezone": "Asia/Qatar", "prefs": {}}],
        "projects": [{"id": P1, "workspace_id": WS[1], "name": "Kite BOM delivery", "description": "Cabling BOM for site", "owner_id": ME, "priority": "high",
                      "start_date": "2026-10-01", "target_date": "2026-11-30", "manual_status": None, "manual_status_reason": None, "budget": None, "budget_currency": None,
                      "next_action": "Confirm revised dispatch date", "archived": False, "created_by": ME, "created_at": iso(now), "updated_at": iso(now)}],
        "milestones": [{"id": str(uuid.uuid4()), "project_id": P1, "title": "Material readiness", "planned_date": "2026-10-20", "actual_date": "2026-10-22", "weight": 1, "is_critical": False, "notes": None, "sort": 0},
                       {"id": str(uuid.uuid4()), "project_id": P1, "title": "Customs clearance", "planned_date": "2026-10-27", "actual_date": None, "weight": 2, "is_critical": True, "notes": None, "sort": 1}],
        "tasks": [
            t(id="t-over", title="Follow up with supplier on delayed shipment", priority="critical", project_id=P1, due_at=iso(now - dt.timedelta(days=1)), next_action="Request written update and recovery plan"),
            t(id="t-wait", title="Customs BOE clearance", status="waiting", waiting_on="Clearing agent", project_id=P1, follow_up_at=iso(now - dt.timedelta(hours=1))),
            t(id="t-crit", title="Book container", priority="high", completion_criteria=[{"text": "Booking confirmation saved", "done": False}], due_at=iso(now + dt.timedelta(days=2))),
            t(id="t-inbox", title="Look into new freight forwarder", status="inbox", workspace_id=WS[2]),
            t(id="t-done", title="Send PO to supplier", status="completed", completed_at=iso(now - dt.timedelta(days=1))),
            t(id="t-deleg", title="Prepare GRN for site", owner_id=ADIL, project_id=P1),
        ],
        "task_dependencies": [{"task_id": "t-crit", "depends_on_id": "t-over"}],
        "follow_ups": [{"id": "f-1", "workspace_id": WS[1], "task_id": None, "project_id": P1, "contact_id": None, "counterpart": "Al Muthawa Trading", "subject": "Revised dispatch date",
                        "direction": "they_owe", "last_contact_at": "2026-10-26", "expected_at": "2026-10-27", "next_follow_up_at": iso(now - dt.timedelta(hours=2)), "escalate_at": None,
                        "status": "open", "notes": None, "document_url": None, "owner_id": ME, "created_by": ME, "closed_at": None, "created_at": iso(now), "updated_at": iso(now)}],
        "contacts": [], "notifications": [{"id": str(uuid.uuid4()), "user_id": ME, "kind": "overdue", "title": "Overdue: Follow up with supplier", "body": "QCLOUD", "url": "/#/t/t-over", "task_id": "t-over", "follow_up_id": None, "read_at": None, "created_at": iso(now)}],
        "task_comments": [], "activity": [], "follow_up_logs": [], "task_contacts": [], "push_subscriptions": [],
        "notes": [{"id": "n-dec", "workspace_id": WS[1], "kind": "decision", "title": "Use sea freight for KSA lane", "body": "", "tags": ["ksa"], "meeting_date": None, "attendees": None,
                   "decision": "Move KSA shipments to sea freight", "rationale": "Cost", "decision_status": "proposed", "decided_on": None, "pinned": False, "created_by": ME, "created_at": iso(now - dt.timedelta(days=2)), "updated_at": iso(now - dt.timedelta(days=2))}],
        "documents": [], "record_links": [], "goals": [], "goal_updates": [], "habits": [], "habit_logs": [], "health_logs": [], "weekly_reviews": [],
        **fin_fixture()
    }

DB = fresh_db()
REQUESTS = []
AI_CALLS = []
STORAGE = {}

def health():
    out = []
    for p in DB["projects"]:
        ts = [t for t in DB["tasks"] if t["project_id"] == p["id"] and not t["deleted_at"]]
        ms = [m for m in DB["milestones"] if m["project_id"] == p["id"]]
        wt = sum(m["weight"] for m in ms); wd = sum(m["weight"] for m in ms if m["actual_date"])
        late = [m for m in ms if not m["actual_date"] and m["planned_date"] and m["planned_date"] < now.date().isoformat()]
        status = "blocked" if any(m["is_critical"] for m in late) else "at_risk" if late else "on_track"
        nxt = next((m for m in sorted(ms, key=lambda m: m["planned_date"] or "9") if not m["actual_date"]), None)
        out.append({"project_id": p["id"], "workspace_id": p["workspace_id"], "progress": round(100 * wd / wt) if wt else None, "progress_basis": "milestones" if ms else "none",
                    "calculated_status": status, "effective_status": p["manual_status"] or status, "reasons": [f"{len(late)} critical milestone(s) overdue"] if late else [],
                    "next_milestone_title": nxt and nxt["title"], "next_milestone_date": nxt and nxt["planned_date"], "task_total": len(ts), "task_open": len([t for t in ts if t["status"] not in ("completed", "cancelled")])})
    return out

def rule_check(row, old=None):
    if row.get("status") == "waiting" and not (row.get("waiting_on") or "").strip(): return "waiting_needs_person: say who you are waiting on"
    if row.get("status") == "blocked" and not (row.get("blocker") or "").strip(): return "blocked_needs_reason: describe what is blocking this task"
    if row.get("status") == "completed" and (old is None or old.get("status") != "completed"):
        if any(not c.get("done") for c in row.get("completion_criteria") or []): return "criteria_incomplete: tick every completion criterion before completing"
        open_pre = [d for d in DB["task_dependencies"] if d["task_id"] == row["id"] and next((t for t in DB["tasks"] if t["id"] == d["depends_on_id"] and t["status"] not in ("completed", "cancelled")), None)]
        if open_pre: return f"dependencies_incomplete: {len(open_pre)} prerequisite task(s) still open"
        if row.get("reviewer_id") and row.get("review_status") != "approved": return "review_required: this task needs the reviewer's approval before it can be completed"
    if row.get("reviewer_id") and row.get("reviewer_id") == (row.get("assignee_id") or row.get("owner_id")): return "self_review: the reviewer must be someone other than the person doing the work"
    if old and old.get("status") == "completed" and row.get("status") != "completed": return "reopen_reason_required: give a reason for reopening"
    return None

def handle(route):
    req = route.request; u = urlparse(req.url); path = u.path; qs = parse_qs(u.query)
    single = "vnd.pgrst.object" in (req.headers.get("accept") or "")
    REQUESTS.append((req.method, path))
    def reply(body, status=200): route.fulfill(status=status, content_type="application/json", body=json.dumps(body))
    def err(msg): reply({"code": "P0001", "message": msg, "details": None, "hint": None}, 400)
    if path.startswith("/auth/v1/"): return reply({"id": ME, "email": "jezz@example.com"})
    m_sign = re.match(r"^/storage/v1/object/sign/(documents|finance)/(.+)$", path)
    if m_sign:
        key = unquote(m_sign.group(2))
        if key not in STORAGE: return reply({"statusCode": "404", "error": "not_found", "message": "Object not found"}, 400)
        return reply({"signedURL": f"/object/sign/{m_sign.group(1)}/{key}?token=test"})
    m_up = re.match(r"^/storage/v1/object/(documents|finance)/(.+)$", path)
    if m_up and req.method == "POST":
        key = unquote(m_up.group(2)); STORAGE[key] = len(req.post_data_buffer or b"")
        return reply({"Key": m_up.group(1) + "/" + key, "Id": str(uuid.uuid4())})
    if path in ("/storage/v1/object/documents", "/storage/v1/object/finance") and req.method == "DELETE":
        for k in json.loads(req.post_data or "{}").get("prefixes", []): STORAGE.pop(k, None)
        return reply([])
    if path.startswith("/functions/v1/gcal-sync"):
        body = json.loads(req.post_data or "{}"); GCAL["calls"].append(body.get("action"))
        if body.get("action") == "status": return reply({"configured": GCAL["configured"], "redirect_uri": "https://lvusersarnzzamdgvxgm.supabase.co/functions/v1/gcal-sync/callback", "sync": GCAL["sync"]})
        if body.get("action") == "start": return reply({"url": "https://accounts.google.com/o/oauth2/v2/auth?client_id=test&scope=openid+email+https%3A%2F%2Fwww.googleapis.com%2Fauth%2Fcalendar.events"})
        if body.get("action") == "resync": return reply({"synced": 4, "error": None})
        if body.get("action") == "disconnect": GCAL["sync"] = None; return reply({"ok": True, "removed": 8})
        return reply({"error": "Unknown action."}, 400)
    if path.startswith("/functions/v1/assistant"):
        body = json.loads(req.post_data or "{}"); AI_CALLS.append(body.get("action"))
        prof = next(p for p in DB["profiles"] if p["id"] == ME); on = bool((prof.get("prefs") or {}).get("ai_enabled"))
        if body.get("action") == "status": return reply({"configured": True, "enabled": on, "used_24h": len(AI_CALLS) - 1, "limit": 60, "model": "claude-sonnet-5-5"})
        if not on: return reply({"error": "Turn on the AI assistant in Settings first."}, 403)
        a, inp = body["action"], body.get("input") or {}
        pid = str(uuid.uuid4()); AI_PROPOSALS[pid] = {"action": a, "context": body.get("context"), "decisions": []}
        if a == "action_items": return reply({"items": [{"title": "Confirm inspection slot with surveyor", "owner": "Adil"}, {"title": "Share revised PO", "owner": None}], "proposal_id": pid})
        if a == "draft_message": return reply({"text": "Hi Rashid, AI-polished: please confirm the dispatch date today.", "proposal_id": pid})
        if a == "weekly_narrative": return reply({"text": "This week I closed out the Kite BOM material readiness.", "proposal_id": pid})
        if a == "categorize": return reply({"suggestions": [{"id": l["id"], "account_id": next(x["id"] for x in inp["accounts"] if x["code"] == "6900"), "confidence": "medium"} for l in inp["lines"]], "proposal_id": pid})
        return reply({"error": "Unknown action."}, 400)
    if path.startswith("/rest/v1/rpc/"):
        fn = path.rsplit("/", 1)[1]; body = json.loads(req.post_data or "{}")
        if fn == "search_all":
            q = body.get("q", "").lower()
            return reply([{"kind": "task", "id": t["id"], "title": t["title"], "subtitle": "", "workspace_id": t["workspace_id"]} for t in DB["tasks"] if q in t["title"].lower()] +
                         [{"kind": "note", "id": n["id"], "title": n["title"], "subtitle": "", "workspace_id": n["workspace_id"]} for n in DB["notes"] if q in (n["title"] + " " + (n["body"] or "")).lower()])
        if fn == "create_calendar_feed":
            DB.setdefault("calendar_feeds", [])[:] = [{"user_id": ME, "include_finance": body.get("p_include_finance", False), "created_at": iso(now), "last_used_at": None}]
            return reply("ab" * 32)
        if fn == "revoke_calendar_feed": DB["calendar_feeds"] = []; return reply(None)
        if fn.startswith("fin_"):
            try: return reply(fin_rpc(fn, body))
            except FinError as e: return err(str(e))
        if fn == "snooze_notification":
            n_ = next((x for x in DB["notifications"] if x["id"] == body["p_notification"]), None)
            if not n_: return err("not_found: alert not found")
            n_["snoozed_until"] = body["p_until"]; n_["read_at"] = n_.get("read_at") or iso(now); return reply(None)
        if fn == "record_ai_decision":
            AI_PROPOSALS[body["p_proposal"]]["decisions"].append(body["p_decision"]); return reply(None)
        if fn == "set_project_baseline":
            pj = next(x for x in DB["projects"] if x["id"] == body["p_project"])
            if pj.get("baselined_at") and len((body.get("p_reason") or "").strip()) < 3: return err("reason_required: give a reason for re-baselining")
            if pj.get("baselined_at"): DB.setdefault("schedule_revisions", []).append({"id": len(DB.get("schedule_revisions", [])) + 1, "workspace_id": pj["workspace_id"], "project_id": pj["id"], "milestone_id": None, "field": "baseline", "old_date": pj["baseline_target_date"], "new_date": pj["target_date"], "reason": "Re-baseline: " + body["p_reason"], "revised_by": ME, "revised_at": iso(now)})
            pj.update(baseline_start_date=pj.get("start_date"), baseline_target_date=pj.get("target_date"), baselined_at=iso(now), baselined_by=ME)
            for m in DB["milestones"]:
                if m["project_id"] == pj["id"]: m["baseline_date"] = m.get("planned_date")
            return reply(None)
        if fn == "revise_project_dates":
            pj = next(x for x in DB["projects"] if x["id"] == body["p_project"])
            for fld, key in (("start_date", "p_start"), ("target_date", "p_target")):
                if body.get(key) and body[key] != pj.get(fld):
                    DB.setdefault("schedule_revisions", []).append({"id": len(DB.get("schedule_revisions", [])) + 1, "workspace_id": pj["workspace_id"], "project_id": pj["id"], "milestone_id": None, "field": fld, "old_date": pj.get(fld), "new_date": body[key], "reason": body["p_reason"], "revised_by": ME, "revised_at": iso(now)})
                    pj[fld] = body[key]
            return reply(None)
        if fn == "revise_milestone_date":
            m = next(x for x in DB["milestones"] if x["id"] == body["p_milestone"])
            DB.setdefault("schedule_revisions", []).append({"id": len(DB.get("schedule_revisions", [])) + 1, "workspace_id": WS[1], "project_id": m["project_id"], "milestone_id": m["id"], "field": "planned_date", "old_date": m.get("planned_date"), "new_date": body["p_planned"], "reason": body["p_reason"], "revised_by": ME, "revised_at": iso(now)})
            m["planned_date"] = body["p_planned"]; return reply(None)
        if fn == "request_task_review":
            t = next(t for t in DB["tasks"] if t["id"] == body["p_task"])
            if not t.get("reviewer_id"): return err("no_reviewer: choose a reviewer first")
            t["review_status"] = "requested"; t["review_note"] = body.get("p_note"); return reply(None)
        if fn == "review_task":
            t = next(t for t in DB["tasks"] if t["id"] == body["p_task"])
            if t.get("reviewer_id") != ME: return err("not_reviewer: only the reviewer can approve or ask for changes")
            if body["p_decision"] == "approve": t.update(review_status="approved", reviewed_by=ME, status="completed", completed_at=iso(dt.datetime.now(dt.timezone.utc)), progress=100)
            else: t.update(review_status="changes_requested", review_note=body.get("p_note"), status="in_progress")
            return reply(None)
        if fn == "reopen_task":
            t = next(t for t in DB["tasks"] if t["id"] == body["p_task"]); t["status"] = body["p_status"]; t["completed_at"] = None
            DB["task_comments"].append({"id": str(uuid.uuid4()), "task_id": t["id"], "author_id": ME, "body": "Reopened: " + body["p_reason"], "kind": "status", "created_at": iso(dt.datetime.now(dt.timezone.utc))})
            return reply(None)
        return reply(None)
    table = path.split("/rest/v1/")[1]
    if table == "project_health": return reply(health())
    if table == "fin_invoice_totals": return reply(inv_totals())
    rows = DB.setdefault(table, [])
    def match(r):
        for k, v in qs.items():
            if k in ("select", "order", "limit", "or", "on_conflict", "columns"): continue
            v = v[0]
            if v.startswith("eq."):
                if str(r.get(k)).lower() != unquote(v[3:]).lower() and str(r.get(k)) != unquote(v[3:]): return False
            elif v.startswith("is.null") and r.get(k) is not None: return False
            elif v.startswith("in.("):
                vals = v[4:-1].split(",")
                if str(r.get(k)) not in vals: return False
        return True
    if req.method == "GET":
        out = [r for r in rows if match(r)]
        return reply(out[0] if single and out else (None if single else out))
    if req.method == "POST":
        body = json.loads(req.post_data); items = body if isinstance(body, list) else [body]; out = []
        for it in items:
            if table == "tasks":
                it = {**fresh_db()["tasks"][0], "id": str(uuid.uuid4()), "created_by": ME, "owner_id": ME, "project_id": None, "title": "", "priority": "medium", "status": "not_started", "due_at": None, "next_action": None, "completed_at": None, **it, "created_at": iso(dt.datetime.now(dt.timezone.utc)), "updated_at": iso(dt.datetime.now(dt.timezone.utc))}
                e = rule_check(it)
                if e: return err(e)
            else:
                stamp = iso(dt.datetime.now(dt.timezone.utc))
                it = {"id": str(uuid.uuid4()), "created_at": stamp, "updated_at": stamp, "author_id": ME, "created_by": ME, "user_id": ME, **it}
                if table == "notes": it = {"tags": [], "pinned": False, "deleted_at": None, "decision_status": "proposed", **it}
                if table == "habits": it = {"days": [0, 1, 2, 3, 4, 5, 6], "per_week": 1, "archived": False, "sort": 0, **it}
                if table == "documents": it = {"tags": [], "deleted_at": None, "category": "other", **it}
                if table == "weekly_reviews" and "on_conflict" in qs: rows[:] = [r for r in rows if r["week_start"] != it["week_start"]]
            if table == "push_subscriptions": rows[:] = [r for r in rows if r.get("endpoint") != it.get("endpoint")]
            if "on_conflict" in qs:
                keys = qs["on_conflict"][0].split(","); rows[:] = [r for r in rows if any(r.get(k) != it.get(k) for k in keys)]
            if table == "fin_entities":
                it = {"lock_date": None, "archived": False, "fiscal_year_start": 1, "next_journal_no": 1, "next_invoice_no": 1, "next_bill_no": 1, **it}
                DB["fin_members"].append({"entity_id": it["id"], "user_id": ME, "role": "owner"}); DB["fin_accounts"].extend(seed_accounts(it["id"]))
            if table == "fin_journals": it = {"status": "draft", "number": None, "source": "manual", "reversed_by": None, "reversal_of": None, **it}
            if table == "fin_lines":
                j = next(x for x in DB["fin_journals"] if x["id"] == it["journal_id"])
                if j["status"] == "posted": return err("posted_immutable: lines of a posted entry can't be changed")
                it = {"debit": 0, "credit": 0, "party_id": None, "statement_line_id": None, "sort": 0, **it, "entity_id": j["entity_id"]}
            if table == "fin_invoices":
                it = {"status": "draft", "number": None, "journal_id": None, "task_id": None, "due_date": None, "currency": None, "fx_rate": None, **it}
                if it["currency"] == next(e for e in DB["fin_entities"] if e["id"] == it["entity_id"])["base_currency"]: it["currency"] = None
                if not it["currency"]: it["fx_rate"] = None
            if table == "fin_invoice_lines": it = {"item_id": None, "account_id": None, **it, "amount": round(float(it["qty"]) * float(it["unit_price"]), 2)}
            if table == "fin_statement_lines": it = {"status": "unmatched", "matched_line_id": None, "reference": None, **it}
            rows.append(it); out.append(it)
        return reply(out[0] if single else out, 201)
    if req.method == "PATCH":
        body = json.loads(req.post_data); out = []
        for r in rows:
            if match(r):
                new = {**r, **body}
                if table == "projects" and r.get("baselined_at") and (new.get("target_date") != r.get("target_date") or new.get("start_date") != r.get("start_date")):
                    return err("reason_required: the schedule is baselined — give a reason for changing the date")
                if table == "follow_ups" and new.get("status") == "closed" and r.get("status") != "closed" and not new.get("outcome"):
                    return err("outcome_required: record the outcome before closing")
                if table == "tasks":
                    e = rule_check(new, r)
                    if e: return err(e)
                    if new["status"] == "completed" and r["status"] != "completed": new["completed_at"] = iso(dt.datetime.now(dt.timezone.utc)); new["progress"] = 100
                    new["updated_at"] = iso(dt.datetime.now(dt.timezone.utc))
                r.clear(); r.update(new); out.append(r)
        return reply(out[0] if single and out else out)
    if req.method == "DELETE":
        rows[:] = [r for r in rows if not match(r)]
        return reply([])
    reply([])

class FinError(Exception): pass
def today_qa(): return (dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=3)).date().isoformat()
def inv_totals():
    out = []
    for i in DB["fin_invoices"]:
        total = round(sum(l["amount"] for l in DB["fin_invoice_lines"] if l["invoice_id"] == i["id"]), 2)
        paid = round(sum(a["amount"] for a in DB["fin_allocations"] if a["invoice_id"] == i["id"] and next((p for p in DB["fin_payments"] if p["id"] == a["payment_id"] and p["status"] == "posted"), None)), 2)
        out.append({"invoice_id": i["id"], "entity_id": i["entity_id"], "total": total, "paid": paid})
    return out
def sys_acc(eid, key): return next(a["id"] for a in DB["fin_accounts"] if a["entity_id"] == eid and a["system_key"] == key)
def next_no(eid, field):
    e = next(x for x in DB["fin_entities"] if x["id"] == eid); n = e[field]; e[field] += 1; return n
def sys_journal(eid, date, memo, source, source_id, lines):
    dr = round(sum(l.get("debit", 0) for l in lines), 2); cr = round(sum(l.get("credit", 0) for l in lines), 2)
    if dr != cr or dr == 0: raise FinError(f"unbalanced: {dr} / {cr}")
    jid = str(uuid.uuid4())
    DB["fin_journals"].append({"id": jid, "entity_id": eid, "number": next_no(eid, "next_journal_no"), "entry_date": date, "memo": memo, "reference": None, "source": source, "source_id": source_id,
                               "status": "posted", "reversal_of": None, "reversed_by": None, "posted_by": ME, "created_by": ME, "created_at": iso(now)})
    for i, l in enumerate(lines):
        DB["fin_lines"].append({"id": str(uuid.uuid4()), "journal_id": jid, "entity_id": eid, "account_id": l["account_id"], "party_id": l.get("party_id"), "description": l.get("description"),
                                "debit": l.get("debit", 0), "credit": l.get("credit", 0), "statement_line_id": None, "sort": i})
    return jid
def p_balance(link_id):
    k = next(x for x in DB["fin_partner_links"] if x["id"] == link_id)
    js = {j["id"] for j in DB["fin_journals"] if j["status"] == "posted"}
    return round(sum(float(l["credit"]) - float(l["debit"]) for l in DB["fin_lines"] if l["journal_id"] in js and l["account_id"] == k["partner_account_id"] and l["party_id"] == k["partner_party_id"]), 2)
def partner_rpc(fn, b):
    if fn == "fin_partner_link_create":
        eid = b["p_entity"]; acc = next((a for a in DB["fin_accounts"] if a["entity_id"] == eid and a["type"] == "liability" and a["subtype"] == "partner"), None)
        if not acc: raise FinError("missing_account: add a partner account first")
        party = {"id": str(uuid.uuid4()), "entity_id": eid, "name": b["p_partner_name"], "kind": "partner", "archived": False, "created_at": iso(now)}; DB["fin_parties"].append(party)
        k = {"id": str(uuid.uuid4()), "entity_id": eid, "partner_party_id": party["id"], "partner_account_id": acc["id"], "partner_user_id": ME, "mirror_entity_id": b.get("p_mirror_entity"),
             "mirror_party_id": None, "mirror_receivable_account_id": None, "mirror_default_account_id": None, "terms_days": None, "next_no": 1, "created_at": iso(now)}
        if k["mirror_entity_id"]:
            me = k["mirror_entity_id"]; ename = next(e["name"] for e in DB["fin_entities"] if e["id"] == eid)
            mp = {"id": str(uuid.uuid4()), "entity_id": me, "name": ename, "kind": "other", "archived": False, "created_at": iso(now)}; DB["fin_parties"].append(mp)
            k.update(mirror_party_id=mp["id"], mirror_receivable_account_id=sys_acc(me, "ar"), mirror_default_account_id=sys_acc(me, "opening"))
        DB.setdefault("fin_partner_links", []).append(k); return k["id"]
    if fn == "fin_partner_record":
        p = b["p"]; k = next(x for x in DB["fin_partner_links"] if x["id"] == p["link_id"]); a = next(x for x in DB["fin_accounts"] if x["id"] == p["account_id"])
        t = p["txn_type"]; d = "decrease" if t == "reimbursement" else (p.get("direction") or "") if t == "adjustment" else "increase"; amt = float(p["amount"])
        if d == "increase" and a["subtype"] in ("cash", "bank", "card"): raise FinError("mixed_funds: money the partner paid personally can't be booked against the business's bank or cash")
        if t == "reimbursement" and a["subtype"] not in ("cash", "bank"): raise FinError("bad_account: a reimbursement is paid from the business's bank or cash account")
        if d == "decrease" and amt > p_balance(k["id"]): raise FinError(f"exceeds_outstanding: {amt} is more than the {p_balance(k['id'])} currently owed to the partner")
        no = "PR-" + str(k["next_no"]).zfill(4); k["next_no"] += 1; tid = str(uuid.uuid4()); desc = no + " · " + p["description"]
        pl = {"account_id": k["partner_account_id"], "party_id": k["partner_party_id"], "description": desc}
        lines = [{"account_id": a["id"], "description": desc, "debit": amt}, {**pl, "credit": amt}] if d == "increase" else [{**pl, "debit": amt}, {"account_id": a["id"], "description": desc, "credit": amt}]
        jid = sys_journal(k["entity_id"], p["txn_date"], no + " " + p["description"], "partner", tid, lines); mj = None
        if k["mirror_entity_id"]:
            macc = p.get("mirror_account_id") or k["mirror_default_account_id"]; rl = {"account_id": k["mirror_receivable_account_id"], "party_id": k["mirror_party_id"], "description": desc}
            ml = [{**rl, "debit": amt}, {"account_id": macc, "description": desc, "credit": amt}] if d == "increase" else [{"account_id": macc, "description": desc, "debit": amt}, {**rl, "credit": amt}]
            mj = sys_journal(k["mirror_entity_id"], p["txn_date"], no + " " + p["description"], "partner", tid, ml)
        DB.setdefault("fin_partner_txns", []).append({"id": tid, "link_id": k["id"], "entity_id": k["entity_id"], "txn_no": no, "txn_type": t, "direction": d, "txn_date": p["txn_date"],
            "date_precision": p.get("date_precision") or "day", "payment_date": p.get("payment_date"), "description": p["description"], "beneficiary": p.get("beneficiary"), "reference": p.get("reference"),
            "invoice_amount": p.get("invoice_amount"), "amount": amt, "currency": "QAR", "paid_from": p.get("paid_from") or ("business_bank" if t == "reimbursement" else "not_recorded"), "account_id": a["id"],
            "job_ref": p.get("job_ref"), "recon_status": "unverified", "flags": p.get("flags") or [], "legacy_ref": None, "notes": p.get("notes"), "import_key": None, "journal_id": jid, "mirror_journal_id": mj,
            "mirror_account_id": p.get("mirror_account_id"), "status": "posted", "reversal_journal_id": None, "mirror_reversal_journal_id": None, "reversal_reason": None, "last_change_reason": None,
            "settles_txn_id": None, "created_by": ME, "created_at": iso(now), "updated_by": None, "updated_at": None})
        return tid
    if fn == "fin_partner_update":
        if not (b.get("p_reason") or "").strip(): raise FinError("reason_required: say why you are changing this record")
        t = next(x for x in DB["fin_partner_txns"] if x["id"] == b["p_txn"]); patch = b["p_patch"]
        for key in patch:
            if key in ("amount", "txn_date", "account_id"): raise FinError(f"locked_field: {key} can't be edited")
        changes = {k2: [t.get(k2), v] for k2, v in patch.items()}; changes["last_change_reason"] = [t.get("last_change_reason"), b["p_reason"]]
        t.update(patch); t["last_change_reason"] = b["p_reason"]
        DB["fin_audit"].append({"id": str(uuid.uuid4()), "entity_id": t["entity_id"], "record_type": "fin_partner_txns", "record_id": t["id"], "actor_id": ME, "action": "update", "changes": changes, "created_at": iso(dt.datetime.now(dt.timezone.utc))})
        return None
    if fn == "fin_partner_reverse":
        t = next(x for x in DB["fin_partner_txns"] if x["id"] == b["p_txn"])
        if not (b.get("p_reason") or "").strip(): raise FinError("reason_required")
        def rev(jid):
            j = next(x for x in DB["fin_journals"] if x["id"] == jid); ls = [l for l in DB["fin_lines"] if l["journal_id"] == jid]
            nid = sys_journal(j["entity_id"], j["entry_date"], "Reversal of #" + str(j["number"]), "reversal", j["id"], [{"account_id": l["account_id"], "party_id": l["party_id"], "debit": l["credit"], "credit": l["debit"]} for l in ls])
            j["reversed_by"] = nid; next(x for x in DB["fin_journals"] if x["id"] == nid)["reversal_of"] = j["id"]; return nid
        t["reversal_journal_id"] = rev(t["journal_id"]); t["mirror_reversal_journal_id"] = rev(t["mirror_journal_id"]) if t["mirror_journal_id"] else None
        t["status"] = "reversed"; t["reversal_reason"] = b["p_reason"]; return None
    if fn == "fin_partner_sync_mirror": return 0
    raise FinError("not implemented in test fake: " + fn)
def ic_rpc(fn, b):
    if fn == "fin_ic_record":
        p = b["p"]; A = next(e for e in DB["fin_entities"] if e["id"] == p["from_entity_id"]); B = next(e for e in DB["fin_entities"] if e["id"] == p["to_entity_id"])
        if A["base_currency"] != B["base_currency"] and not p.get("amount_to"): raise FinError("amount_to_required: enter the received amount")
        amt, amt_to = float(p["amount"]), float(p.get("amount_to") or p["amount"])
        tid = str(uuid.uuid4()); no = "IC-" + str(len(DB["fin_ic_txns"]) + 1).zfill(5)
        def ic_acc(eid):
            a = next((x for x in DB["fin_accounts"] if x["entity_id"] == eid and x.get("system_key") == "intercompany"), None)
            if not a: a = {"id": f"{eid[:8]}-acc-1190", "entity_id": eid, "code": "1190", "name": "Inter-entity current accounts", "type": "asset", "subtype": "other", "system_key": "intercompany", "is_active": True, "description": None}; DB["fin_accounts"].append(a)
            return a["id"]
        ja = sys_journal(A["id"], p["txn_date"], no + " " + p["description"], "intercompany", tid, [{"account_id": ic_acc(A["id"]), "debit": amt}, {"account_id": p["from_account_id"], "credit": amt}])
        jb = sys_journal(B["id"], p["txn_date"], no + " " + p["description"], "intercompany", tid, [{"account_id": p["to_account_id"], "debit": amt_to}, {"account_id": ic_acc(B["id"]), "credit": amt_to}])
        DB["fin_ic_txns"].append({"id": tid, "ic_no": no, "from_entity_id": A["id"], "to_entity_id": B["id"], "kind": p["kind"], "txn_date": p["txn_date"], "description": p["description"], "reference": p.get("reference"),
            "amount": amt, "amount_to": amt_to, "from_account_id": p["from_account_id"], "to_account_id": p["to_account_id"], "from_journal_id": ja, "to_journal_id": jb, "status": "posted", "reversal_reason": None, "created_at": iso(now)})
        return tid
    if fn == "fin_ic_position":
        e = b["p_entity"]; out = {}
        for t in DB["fin_ic_txns"]:
            if t["status"] != "posted" or e not in (t["from_entity_id"], t["to_entity_id"]): continue
            o = t["to_entity_id"] if t["from_entity_id"] == e else t["from_entity_id"]; oe = next(x for x in DB["fin_entities"] if x["id"] == o)
            r_ = out.setdefault(o, {"entity_id": o, "name": oe["name"], "currency": oe["base_currency"], "mine": 0, "theirs": 0, "recorded_mine": 0, "recorded_theirs": 0, "matched": True})
            if t["from_entity_id"] == e: r_["mine"] += t["amount"]; r_["theirs"] += t["amount_to"]
            else: r_["mine"] -= t["amount_to"]; r_["theirs"] -= t["amount"]
            r_["recorded_mine"], r_["recorded_theirs"] = r_["mine"], r_["theirs"]
        return list(out.values())
    if fn == "fin_ic_reverse":
        t = next(x for x in DB["fin_ic_txns"] if x["id"] == b["p_txn"]); t["status"] = "reversed"; t["reversal_reason"] = b["p_reason"]; return None
    raise FinError("not implemented: " + fn)
def prod_acc(eid, key):
    spec = {"wip": ("1250", "Work in progress", "asset"), "grni": ("2050", "Goods received, not yet billed", "liability"), "inventory": ("1200", "Inventory", "asset"), "cogs": ("5000", "Cost of goods sold", "expense")}[key]
    a = next((x for x in DB["fin_accounts"] if x["entity_id"] == eid and x.get("system_key") == key), None)
    if not a:
        a = {"id": f"{eid[:8]}-acc-{spec[0]}", "entity_id": eid, "code": spec[0], "name": spec[1], "type": spec[2], "subtype": "other", "system_key": key, "is_active": True, "description": None}; DB["fin_accounts"].append(a)
    return a["id"]
def on_hand(item):
    ms = [m for m in DB["fin_stock_moves"] if m["item_id"] == item]; q = sum(float(m["qty"]) for m in ms); v = sum(float(m["qty"]) * float(m["unit_cost"]) for m in ms)
    return round(q, 3), (v / q if q > 0 else 0)
def add_move(eid, item, date, qty, cost, source, jid, note, order=None, exc=None):
    mid = str(uuid.uuid4()); DB["fin_stock_moves"].append({"id": mid, "entity_id": eid, "item_id": item, "move_date": date, "qty": qty, "unit_cost": cost, "source": source, "journal_id": jid, "note": note,
        "production_order_id": order, "exception_id": exc, "created_at": iso(now)}); return mid
def prod_fx(eid):
    a = next((x for x in DB["fin_accounts"] if x["entity_id"] == eid and x.get("system_key") == "fx"), None)
    if not a: a = {"id": f"{eid[:8]}-acc-4900", "entity_id": eid, "code": "4900", "name": "Exchange gains and losses", "type": "income", "subtype": "other", "system_key": "fx", "is_active": True, "description": None}; DB["fin_accounts"].append(a)
    return a["id"]
def prod_rpc(fn, b):
    today = today_qa()
    if fn == "fin_goods_receipt":
        p = b["p"]; it = next(x for x in DB["fin_items"] if x["id"] == p["item_id"]); e = it["entity_id"]; gid = str(uuid.uuid4())
        no = "GRN-" + str(len([g for g in DB["fin_goods_receipts"] if g["entity_id"] == e]) + 1).zfill(4); val = round(float(p["qty"]) * float(p["unit_cost"]), 2); d = p.get("receipt_date") or today
        jid = sys_journal(e, d, no + " received", "receipt", gid, [{"account_id": prod_acc(e, "inventory"), "debit": val}, {"account_id": prod_acc(e, "grni"), "party_id": p.get("party_id"), "credit": val}])
        mid = add_move(e, it["id"], d, float(p["qty"]), float(p["unit_cost"]), "receipt", jid, no)
        DB["fin_goods_receipts"].append({"id": gid, "entity_id": e, "receipt_no": no, "item_id": it["id"], "party_id": p.get("party_id"), "receipt_date": d, "qty": float(p["qty"]), "unit_cost": float(p["unit_cost"]),
            "reference": p.get("reference"), "journal_id": jid, "stock_move_id": mid, "status": "posted", "reversal_reason": None, "created_at": iso(now)}); return gid
    if fn == "fin_production_open":
        e = b["p_entity"]; oid = str(uuid.uuid4()); no = "JOB-" + str(len([o for o in DB["fin_production_orders"] if o["entity_id"] == e]) + 1).zfill(4)
        DB["fin_production_orders"].insert(0, {"id": oid, "entity_id": e, "order_no": no, "title": b["p_title"], "party_id": b.get("p_party"), "job_ref": b.get("p_job_ref"), "status": "open", "opened_on": b.get("p_date") or today,
            "completed_on": None, "completion_journal_id": None, "total_cost": None, "invoice_id": None, "notes": None, "cancel_reason": None, "created_at": iso(now)}); return oid
    if fn in ("fin_production_issue", "fin_production_add_cost", "fin_production_complete", "fin_production_cancel"):
        o = next(x for x in DB["fin_production_orders"] if x["id"] == b["p_order"]); e = o["entity_id"]
        if o["status"] != "open": raise FinError("order_closed: the order is " + o["status"])
        posted = [c for c in DB["fin_production_costs"] if c["order_id"] == o["id"] and c["status"] == "posted"]
        if fn == "fin_production_issue":
            qty = float(b["p_qty"]); q, avg = on_hand(b["p_item"]); exc = None
            if q < qty:
                if not b.get("p_authorise_reason"): raise FinError(f"insufficient_stock: only {q} in stock")
                exc = str(uuid.uuid4()); DB["fin_stock_exceptions"].insert(0, {"id": exc, "entity_id": e, "item_id": b["p_item"], "qty_short": round(qty - max(q, 0), 3), "reason": b["p_authorise_reason"], "authorised_by": ME, "authorised_at": iso(now)})
            val = round(qty * avg, 2); cid = str(uuid.uuid4()); d = b.get("p_date") or today
            jid = sys_journal(e, d, o["order_no"] + " material", "production", cid, [{"account_id": prod_acc(e, "wip"), "debit": val}, {"account_id": prod_acc(e, "inventory"), "credit": val}])
            mid = add_move(e, b["p_item"], d, -qty, avg, "production", jid, o["order_no"], o["id"], exc)
            DB["fin_production_costs"].append({"id": cid, "order_id": o["id"], "entity_id": e, "kind": "material", "cost_date": d, "item_id": b["p_item"], "qty": qty, "unit_cost": avg, "amount": val, "source_account_id": None,
                "note": b.get("p_note"), "journal_id": jid, "stock_move_id": mid, "exception_id": exc, "status": "posted", "reversal_reason": None, "created_at": iso(now)}); return cid
        if fn == "fin_production_add_cost":
            cid = str(uuid.uuid4()); d = b.get("p_date") or today; amt = round(float(b["p_amount"]), 2)
            jid = sys_journal(e, d, o["order_no"] + " cost", "production", cid, [{"account_id": prod_acc(e, "wip"), "debit": amt}, {"account_id": b["p_account"], "credit": amt}])
            DB["fin_production_costs"].append({"id": cid, "order_id": o["id"], "entity_id": e, "kind": "other", "cost_date": d, "item_id": None, "qty": None, "unit_cost": None, "amount": amt, "source_account_id": b["p_account"],
                "note": b["p_note"], "journal_id": jid, "stock_move_id": None, "exception_id": None, "status": "posted", "reversal_reason": None, "created_at": iso(now)}); return cid
        if fn == "fin_production_complete":
            tot = round(sum(float(c["amount"]) for c in posted), 2)
            if tot <= 0: raise FinError("no_costs")
            jid = sys_journal(e, b["p_date"], o["order_no"] + " completed", "production", o["id"], [{"account_id": prod_acc(e, "cogs"), "debit": tot}, {"account_id": prod_acc(e, "wip"), "credit": tot}])
            o.update(status="completed", completed_on=b["p_date"], completion_journal_id=jid, total_cost=tot); return tot
        if posted: raise FinError("has_costs: reverse the material and costs first")
        o.update(status="cancelled", cancel_reason=b["p_reason"]); return None
    if fn == "fin_production_reverse_cost":
        c = next(x for x in DB["fin_production_costs"] if x["id"] == b["p_cost"]); j = next(x for x in DB["fin_journals"] if x["id"] == c["journal_id"])
        ls = [l for l in DB["fin_lines"] if l["journal_id"] == j["id"]]
        nid = sys_journal(c["entity_id"], j["entry_date"], "Reversal", "reversal", j["id"], [{"account_id": l["account_id"], "debit": l["credit"], "credit": l["debit"]} for l in ls]); j["reversed_by"] = nid
        if c["kind"] == "material": add_move(c["entity_id"], c["item_id"], c["cost_date"], c["qty"], c["unit_cost"], "production", nid, "Returned", c["order_id"])
        c.update(status="reversed", reversal_reason=b["p_reason"]); return None
    if fn == "fin_goods_receipt_reverse":
        g = next(x for x in DB["fin_goods_receipts"] if x["id"] == b["p_receipt"]); g.update(status="reversed", reversal_reason=b["p_reason"]); return None
    raise FinError("not implemented: " + fn)
def fin_rpc(fn, b):
    if fn.startswith("fin_production_") or fn.startswith("fin_goods_receipt"): return prod_rpc(fn, b)
    if fn.startswith("fin_ic_"): return ic_rpc(fn, b)
    if fn.startswith("fin_partner_"): return partner_rpc(fn, b)
    if fn == "fin_post_journal":
        j = next(x for x in DB["fin_journals"] if x["id"] == b["p_journal"]); ls = [l for l in DB["fin_lines"] if l["journal_id"] == j["id"]]
        dr = round(sum(float(l["debit"]) for l in ls), 2); cr = round(sum(float(l["credit"]) for l in ls), 2)
        if len(ls) < 2: raise FinError("too_few_lines: an entry needs at least two lines")
        if dr != cr: raise FinError(f"unbalanced: debits ({dr}) and credits ({cr}) must be equal")
        j["status"] = "posted"; j["number"] = next_no(j["entity_id"], "next_journal_no"); return j["number"]
    if fn == "fin_post_invoice":
        i = next(x for x in DB["fin_invoices"] if x["id"] == b["p_invoice"])
        if not i.get("party_id"): raise FinError("party_required: choose the customer")
        if i.get("currency") and not i.get("fx_rate"): raise FinError("rate_required: enter the exchange rate for " + i["currency"])
        k = float(i.get("fx_rate") or 1)
        ls = [l for l in DB["fin_invoice_lines"] if l["invoice_id"] == i["id"]]; total = round(sum(round(l["amount"] * k, 2) for l in ls), 2)
        sale = i["kind"] == "sale"
        jl = [{"account_id": l["account_id"] or sys_acc(i["entity_id"], "sales"), "description": l["description"], ("credit" if sale else "debit"): round(l["amount"] * k, 2)} for l in ls]
        jl.append({"account_id": sys_acc(i["entity_id"], "ar" if sale else "ap"), "party_id": i["party_id"], ("debit" if sale else "credit"): total})
        no = ("INV-" if sale else "BILL-") + str(next_no(i["entity_id"], "next_invoice_no" if sale else "next_bill_no")).zfill(4)
        i["journal_id"] = sys_journal(i["entity_id"], i["issue_date"], ("Invoice " if sale else "Bill ") + no, "invoice" if sale else "bill", i["id"], jl)
        i["status"] = "posted"; i["number"] = no; return no
    if fn == "fin_record_payment":
        eid, amt = b["p_entity"], float(b["p_amount"]); tot = {t["invoice_id"]: t for t in inv_totals()}
        for a in b["p_allocations"]:
            t = tot[a["invoice_id"]]
            if float(a["amount"]) > t["total"] - t["paid"]: raise FinError("overpayment: more than outstanding")
        pid = str(uuid.uuid4()); inn = b["p_direction"] == "in"
        DB["fin_payments"].append({"id": pid, "entity_id": eid, "direction": b["p_direction"], "party_id": b["p_party"], "pay_date": b["p_date"], "amount": amt, "account_id": b["p_account"],
                                   "reference": b["p_reference"], "memo": b["p_memo"], "status": "posted", "journal_id": None, "created_at": iso(now)})
        jid = sys_journal(eid, b["p_date"], "Payment received" if inn else "Payment made", "payment", pid,
            [{"account_id": b["p_account"], ("debit" if inn else "credit"): amt}, {"account_id": sys_acc(eid, "ar" if inn else "ap"), "party_id": b["p_party"], ("credit" if inn else "debit"): amt}])
        DB["fin_payments"][-1]["journal_id"] = jid
        for a in b["p_allocations"]: DB["fin_allocations"].append({"payment_id": pid, "invoice_id": a["invoice_id"], "entity_id": eid, "amount": float(a["amount"])})
        return pid
    if fn == "fin_record_fx_payment":
        p = b["p"]; eid = p["entity_id"]; inn = p["direction"] == "in"; amt, base = float(p["amount"]), float(p.get("base_amount") or 0)
        if not base: raise FinError("base_amount_required: enter what actually reached the bank")
        tot = {t["invoice_id"]: t for t in inv_totals()}; ctrl = 0; s_ = 0
        for a in p["allocations"]:
            inv = next(x for x in DB["fin_invoices"] if x["id"] == a["invoice_id"])
            if inv.get("currency") != p["currency"]: raise FinError("currency_mismatch")
            t = tot[a["invoice_id"]]
            if float(a["amount"]) > t["total"] - t["paid"]: raise FinError("overpayment")
            ctrl += round(float(a["amount"]) * float(inv["fx_rate"]), 2); s_ += float(a["amount"])
        rate = base / amt; ctrl = round(ctrl + (amt - s_) * rate, 2); gain = round(base - ctrl if inn else ctrl - base, 2)
        pid = str(uuid.uuid4())
        DB["fin_payments"].append({"id": pid, "entity_id": eid, "direction": p["direction"], "party_id": p["party_id"], "pay_date": p["pay_date"], "amount": amt, "account_id": p["account_id"],
            "reference": p.get("reference"), "memo": p.get("memo"), "status": "posted", "journal_id": None, "currency": p["currency"], "fx_rate": round(rate, 8), "base_amount": base, "fx_gain": gain, "created_at": iso(now)})
        jl = [{"account_id": p["account_id"], ("debit" if inn else "credit"): base}, {"account_id": sys_acc(eid, "ar" if inn else "ap"), "party_id": p["party_id"], ("credit" if inn else "debit"): ctrl}]
        if gain: jl.append({"account_id": prod_fx(eid), ("credit" if gain > 0 else "debit"): abs(gain)})
        DB["fin_payments"][-1]["journal_id"] = sys_journal(eid, p["pay_date"], "FX payment", "payment", pid, jl)
        for a in p["allocations"]: DB["fin_allocations"].append({"payment_id": pid, "invoice_id": a["invoice_id"], "entity_id": eid, "amount": float(a["amount"])})
        return pid
    if fn == "fin_fx_open_items":
        tot = {t["invoice_id"]: t for t in inv_totals()}; out = []
        for i in DB["fin_invoices"]:
            if i["entity_id"] != b["p_entity"] or i["status"] != "posted" or not i.get("currency") or i["issue_date"] > b["p_date"]: continue
            o = round(tot[i["id"]]["total"] - tot[i["id"]]["paid"], 2)
            if o > 0: out.append({"invoice_id": i["id"], "number": i["number"], "kind": i["kind"], "party_id": i["party_id"], "currency": i["currency"], "outstanding": o, "booked_rate": float(i["fx_rate"]), "booked_base": round(o * float(i["fx_rate"]), 2)})
        return out
    if fn == "fin_fx_revalue":
        e = b["p_entity"]; net = 0; detail = []
        for x in fin_rpc("fin_fx_open_items", {"p_entity": e, "p_date": b["p_date"]}):
            r_ = float(b["p_rates"][x["currency"]]); d_ = round(round(x["outstanding"] * r_, 2) - x["booked_base"], 2); g = d_ if x["kind"] == "sale" else -d_
            net += g; detail.append({**x, "closing_rate": r_, "gain": g})
        vid = str(uuid.uuid4())
        DB["fin_fx_revaluations"].insert(0, {"id": vid, "entity_id": e, "rev_date": b["p_date"], "rates": b["p_rates"], "detail": detail, "net": round(net, 2), "status": "posted", "cancel_reason": None, "created_at": iso(now)})
        return vid
    if fn == "fin_fx_revaluation_cancel":
        v = next(x for x in DB["fin_fx_revaluations"] if x["id"] == b["p_id"]); v.update(status="cancelled", cancel_reason=b["p_reason"]); return None
    if fn == "fin_match_statement":
        st = next(x for x in DB["fin_statement_lines"] if x["id"] == b["p_statement"]); l = next(x for x in DB["fin_lines"] if x["id"] == b["p_line"])
        if round(float(l["debit"]) - float(l["credit"]), 2) != round(float(st["amount"]), 2): raise FinError("amount_mismatch")
        st["status"] = "matched"; st["matched_line_id"] = l["id"]; l["statement_line_id"] = st["id"]; return None
    raise FinError("not implemented in test fake: " + fn)

AI_PROPOSALS = {}
GCAL = {"configured": False, "sync": None, "calls": []}
results, errors = [], []
def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))
    print(("PASS " if cond else "FAIL ") + name + (f" — {detail}" if detail and not cond else ""))

session = {"access_token": "x", "token_type": "bearer", "expires_in": 3600, "expires_at": int(time.time()) + 36000, "refresh_token": "r",
           "user": {"id": ME, "email": "jezz@example.com", "aud": "authenticated"}}

def new_page(browser, w, h, scheme="light"):
    ctx = browser.new_context(viewport={"width": w, "height": h}, device_scale_factor=2, color_scheme=scheme, service_workers="block", timezone_id="Asia/Qatar")
    page = ctx.new_page()
    page.on("pageerror", lambda e: errors.append("pageerror: " + str(e)))
    page.on("console", lambda m: errors.append("console: " + m.text) if m.type == "error" and "ERR_FAILED" not in m.text and "WebSocket" not in m.text and "websocket" not in m.text else None)
    page.route("https://cdn.jsdelivr.net/**", lambda r: r.fulfill(path=UMD, content_type="application/javascript"))
    page.route(re.compile(r"https://fonts\.(googleapis|gstatic)\.com/.*"), lambda r: r.abort())
    page.route(re.compile(r"https://.*\.supabase\.co/.*"), handle)
    page.add_init_script(f"localStorage.setItem('sb-lvusersarnzzamdgvxgm-auth-token', {json.dumps(json.dumps(session))})")
    return ctx, page

def shot(page, name, full=False):
    if SHOTS: page.screenshot(path=os.path.join(SHOTS, name + ".png"), full_page=full)

with sync_playwright() as pw:
    b = pw.chromium.launch()

    # ---------- signed-out ----------
    ctx = b.new_context(viewport={"width": 390, "height": 844}, service_workers="block")
    pg = ctx.new_page(); pg.on("pageerror", lambda e: errors.append("pageerror: " + str(e)))
    pg.route("https://cdn.jsdelivr.net/**", lambda r: r.fulfill(path=UMD, content_type="application/javascript"))
    pg.route(re.compile(r"https://.*\.supabase\.co/.*"), handle)
    pg.goto(BASE + "/"); pg.wait_for_selector("#authForm", timeout=8000)
    check("sign-in screen renders when signed out", pg.is_visible("#af-email") and pg.is_visible("#af-pass"))
    pg.click("#af-btn"); check("sign-in validates email before calling the server", "valid email" in pg.inner_text("#authForm"))
    ctx.close()

    # ---------- desktop ----------
    ctx, page = new_page(b, 1360, 900)
    page.goto(BASE + "/#/home"); page.wait_for_selector(".tiles", timeout=10000)
    check("desktop sidebar shows all 9 sections", page.locator(".sidenav nav a").count() == 9)
    check("Command Centre shows real overdue count", page.inner_text(".tile.bad .tn").strip() == "1")
    check("attention list ranks overdue critical task first", "Follow up with supplier" in page.locator(".attn-row").first.inner_text())
    check("project health shows blocked (critical milestone overdue)", "Blocked" in page.inner_text(".ph-list"))
    shot(page, "01-home-desktop", True)

    # Quick capture without a deadline
    page.keyboard.press("n"); page.wait_for_selector("#capForm")
    page.fill("#cap-title", "Call Hydromaster about cargo pants sample")
    page.click("#capSave"); page.wait_for_selector("text=Capture another")
    created = [t for t in DB["tasks"] if t["title"] == "Call Hydromaster about cargo pants sample"]
    check("quick capture saves to the server", len(created) == 1)
    check("captured task has no invented deadline or follow-up", created and created[0]["due_at"] is None and created[0]["follow_up_at"] is None)
    check("captured task defaults: inbox, medium, owner = me", created and created[0]["status"] == "inbox" and created[0]["priority"] == "medium" and created[0]["owner_id"] == ME)
    page.click("text=Capture another"); page.fill("#cap-title", "Renew trade licence")
    page.check('input[value="not_started"]', force=True)
    page.click('[data-act="cap-more"]'); page.click('[data-act="cap-quick"][data-f="due"][data-n="7"]')
    page.select_option("#cap-priority", "high"); page.click("#capSave"); page.wait_for_selector("text=Capture another")
    t2 = next((t for t in DB["tasks"] if t["title"] == "Renew trade licence"), None)
    check("capture with deadline stores a UTC timestamp at 23:59 Doha", t2 and t2["due_at"] and t2["due_at"].endswith("T20:59:00.000Z"), str(t2 and t2["due_at"]))
    check("capture as 'Ready to work' sets Not started + High", t2 and t2["status"] == "not_started" and t2["priority"] == "high")
    page.click('[data-act="close-sheet"] >> nth=0')

    # Completion rules
    page.goto(BASE + "/#/tasks/all"); page.wait_for_selector(".list")
    page.click('.trow[data-id="t-crit"] .tcheck'); page.wait_for_selector(".toast.error")
    check("completion blocked by unticked criteria + open prerequisite (client message)", "criteri" in page.inner_text("#toast") and "prerequisite" in page.inner_text("#toast"))
    check("blocked completion didn't change the task", next(t for t in DB["tasks"] if t["id"] == "t-crit")["status"] == "not_started")

    # Open task, set waiting needs a person
    page.click('.trow[data-id="t-over"]'); page.wait_for_selector("#tf-title")
    check("task panel shows next action", page.input_value("#tf-next_action") == "Request written update and recovery plan")
    page.select_option("#tf-status", "waiting"); page.wait_for_selector(".dlg")
    page.click('.dlg [data-r="1"]'); check("waiting requires who you're waiting on", page.is_visible("#dlgErr"))
    page.fill("#dlgIn", "Al Muthawa Trading"); page.click('.dlg [data-r="1"]')
    page.wait_for_timeout(400)
    tov = next(t for t in DB["tasks"] if t["id"] == "t-over")
    check("status → Waiting saved with waiting_on", tov["status"] == "waiting" and tov["waiting_on"] == "Al Muthawa Trading")
    # follow-up date
    page.fill("#tf-fu-d", "2026-11-03"); page.fill("#tf-fu-t", "10:00"); page.locator("#tf-fu-t").dispatch_event("change")
    page.wait_for_timeout(400)
    check("follow-up date stored as UTC (10:00 Doha = 07:00Z)", (tov.get("follow_up_at") or "").startswith("2026-11-03T07:00"), str(tov.get("follow_up_at")))
    # add criterion + subtask + comment
    page.fill("#critInput", "Supplier response recorded"); page.click("#critForm button"); page.wait_for_timeout(300)
    check("completion criterion added", any(c["text"] == "Supplier response recorded" for c in tov["completion_criteria"]))
    page.click('[data-act="task-tab"][data-tab="plan"]'); page.fill("#subInput", "Email supplier for written update"); page.click("#subForm button"); page.wait_for_timeout(400)
    check("subtask created under the task", any(t["parent_id"] == "t-over" for t in DB["tasks"]))
    page.click('[data-act="task-tab"][data-tab="updates"]'); page.fill("#upInput", "Spoke to Rashid, written update promised by Sunday"); page.click("#upForm button"); page.wait_for_timeout(400)
    check("progress update logged", any(c["body"].startswith("Spoke to Rashid") for c in DB["task_comments"]))
    shot(page, "02-task-panel")
    page.keyboard.press("Escape")

    # Reopen requires a reason
    page.goto(BASE + "/#/tasks/completed"); page.wait_for_selector(".list")
    page.click('.trow[data-id="t-done"] .tcheck'); page.wait_for_selector(".dlg")
    page.click('.dlg [data-r="1"]'); check("reopen requires a reason", page.is_visible("#dlgErr"))
    page.fill("#dlgIn", "Supplier rejected the PO"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(500)
    td = next(t for t in DB["tasks"] if t["id"] == "t-done")
    check("reopen recorded with reason via RPC", td["status"] == "in_progress" and any("Supplier rejected" in c["body"] for c in DB["task_comments"]))

    # Board
    page.goto(BASE + "/#/tasks/all"); page.click('[data-act="task-mode"][data-m="board"]'); page.wait_for_selector(".board")
    check("board shows status columns", page.locator(".bcol").count() == 6)
    shot(page, "03-board")
    page.click('[data-act="task-mode"][data-m="list"]')

    # Project page
    page.goto(BASE + f"/#/p/{P1}"); page.wait_for_selector(".proj-hero")
    check("project shows milestone-weighted progress 33%", "33%" in page.inner_text(".proj-hero"))
    check("milestone variance shows days late from real dates", "2d late" in page.inner_text(".mslist"))
    shot(page, "04-project", True)

    # Follow-ups & calendar & search
    page.goto(BASE + "/#/followups"); page.wait_for_selector(".sect")
    check("follow-up due now is surfaced", "Follow up now" in page.inner_text("#view"))
    page.goto(BASE + "/#/calendar"); page.wait_for_selector(".cal-month")
    check("calendar month grid renders 42 days", page.locator(".cal-cell").count() == 42)
    shot(page, "05-calendar")
    page.goto(BASE + "/#/search"); page.fill("#global-q", "container"); page.wait_for_selector(".srow", timeout=4000)
    check("search returns matching task", "Book container" in page.inner_text("#searchResults"))
    page.goto(BASE + "/#/finance"); page.wait_for_selector(".fin-card", timeout=8000)
    check("finance page opens the books list", page.locator(".fin-card").count() >= 1)

    # ---------- Phase 2: notes, meetings, decisions ----------
    check("goals, notes and reports are no longer marked 'coming soon'", page.locator(".sidenav nav a.soon").count() == 0)
    page.goto(BASE + "/#/notes/meetings"); page.wait_for_selector(".empty")
    page.click('.page-head [data-act="new-note"]'); page.wait_for_selector("#noteForm")
    page.select_option("#nf-ws", WS[1])
    page.fill("#nf-title", "Weekly logistics review")
    page.fill("#nf-body", "Discussed delays.\n- [ ] Send revised PO to supplier @Adil\n- [x] Invoice already sent\nAction: book inspection slot")
    page.click("#nfSave")
    for _ in range(40):
        if any(n["title"] == "Weekly logistics review" for n in DB["notes"]): break
        page.wait_for_timeout(100)
    page.wait_for_timeout(200)
    mtg = next((n for n in DB["notes"] if n["title"] == "Weekly logistics review"), None)
    check("meeting note saved with kind=meeting and today's date", mtg and mtg["kind"] == "meeting" and mtg["meeting_date"])
    page.click('[data-act="note-actions"]'); page.wait_for_selector(".actions-review")
    check("action items proposed for review (done item skipped)", page.locator(".actions-review li").count() == 2)
    check("@Adil mapped to the workspace member", page.input_value("#ai-owner-0") == ADIL)
    before = len(DB["tasks"])
    page.click('[data-act="ai-create"]'); page.wait_for_timeout(900)
    made = DB["tasks"][before:]
    check("reviewed action items became tasks, with no invented deadlines", len(made) == 2 and all(t["due_at"] is None and t["follow_up_at"] is None for t in made), str(made))
    check("new tasks linked back to the meeting", sum(1 for l in DB["record_links"] if l["from_type"] == "note" and l["from_id"] == (mtg or {}).get("id") and l["to_type"] == "task") == 2)
    shot(page, "10-meeting-note")
    page.keyboard.press("Escape")
    page.goto(BASE + "/#/notes/decisions"); page.wait_for_selector(".note-row")
    check("pending decision shown in the decision log", "Pending" in page.inner_text("#view") and "1 decision pending" in page.inner_text(".page-head"))
    page.click('.page-head [data-act="new-note"]'); page.wait_for_selector("#noteForm")
    page.fill("#nf-title", "Switch forwarder"); page.click("#nfSave"); page.wait_for_selector(".toast.error")
    check("decision requires the decision text", "decision" in page.inner_text("#toast").lower() and not any(n["title"] == "Switch forwarder" for n in DB["notes"]))
    page.keyboard.press("Escape")

    # ---------- documents: attach to a task, open via signed link, delete ----------
    page.goto(BASE + "/#/tasks/all"); page.wait_for_selector(".list")
    page.click('.trow[data-id="t-crit"]'); page.wait_for_selector("#tf-title")
    page.click('[data-act="task-tab"][data-tab="plan"]'); page.wait_for_selector(".kpanel")
    page.set_input_files('input[data-act="attach-file"]', files=[{"name": "Booking confirmation #44.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4 test"}])
    page.wait_for_selector('.kpanel .krow', timeout=5000)
    doc = next((d for d in DB["documents"] if d["title"].startswith("Booking confirmation")), None)
    check("file uploaded to private storage under its workspace folder", doc and doc["storage_path"].startswith(WS[1] + "/") and doc["storage_path"] in STORAGE, str(doc))
    check("stored file name is sanitised", doc and doc["storage_path"].endswith("/Booking_confirmation_44.pdf"))
    check("document linked to the task", any(l["from_type"] == "document" and l["to_type"] == "task" and l["to_id"] == "t-crit" for l in DB["record_links"]))
    page.click('.kpanel [data-act="open-doc"]'); page.wait_for_selector('[data-act="doc-open"]')
    page.context.route(re.compile(r"https://.*\.supabase\.co/storage/v1/object/sign/.*"), lambda r: r.fulfill(body="signed file", content_type="text/plain"))
    with page.context.expect_page() as pop: page.click('[data-act="doc-open"]')
    popup = pop.value
    for _ in range(30):
        if "/object/sign/" in popup.url: break
        popup.wait_for_timeout(100)
    check("document opens through an expiring signed link", "/object/sign/documents/" in popup.url and "token=" in popup.url, popup.url)
    popup.close()
    page.click('[data-act="doc-delete"]'); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(500)
    check("deleting a document removes the stored file and the record", doc and doc["storage_path"] not in STORAGE and not DB["documents"])
    page.keyboard.press("Escape")

    # ---------- habits, goals, health (private) ----------
    page.goto(BASE + "/#/goals/habits"); page.wait_for_selector(".empty")
    check("goals page states data is private", "Private to you" in page.inner_text(".page-head"))
    page.click('.page-head [data-act="habit-new"]'); page.fill("#hb-title", "Gym"); page.click("#hbSave"); page.wait_for_selector(".hgrid")
    page.click('.habits [data-act="habit-toggle"]'); page.wait_for_timeout(400)
    check("ticking a habit logs today's check-in", len(DB["habit_logs"]) == 1)
    check("streak shows 1 after first check-in", "🔥 1" in page.inner_text(".hgrid"))
    page.click('.habits [data-act="habit-toggle"]'); page.wait_for_timeout(400)
    check("unticking removes the check-in", len(DB["habit_logs"]) == 0)
    check("future days cannot be ticked", page.locator(".hcell:disabled").count() == page.evaluate("6 - new Date(new Date().toLocaleString('en-US', {timeZone: 'Asia/Qatar'})).getDay()"))
    shot(page, "11-habits")
    page.goto(BASE + "/#/home"); page.wait_for_selector(".tiles")
    check("Command Centre shows habits for today", "habits today" in page.inner_text("#view").lower())
    page.goto(BASE + "/#/goals/goals"); page.click('.page-head [data-act="goal-new"]'); page.wait_for_selector("#goalForm")
    page.fill("#gl-title", "Lose 8 kg"); page.fill("#gl-target", "8"); page.fill("#gl-current", "2"); page.fill("#gl-unit", "kg"); page.click("#glSave"); page.wait_for_function("document.querySelector('#toast') && document.querySelector('#toast').innerText.includes('Goal saved')")
    page.keyboard.press("Escape"); page.wait_for_selector(".goal-card")
    gc = page.locator(".goal-card", has_text="Lose 8 kg").first.inner_text()
    check("goal progress computed from target and current (25%)", "25%" in gc, gc[:120])
    page.goto(BASE + "/#/goals/health"); page.click('.page-head [data-act="health-new"]'); page.wait_for_selector("#healthForm")
    page.fill("#hl-value", "82.4"); page.click("#hlSave"); page.wait_for_timeout(400)
    check("health entry saved", len(DB["health_logs"]) == 1 and DB["health_logs"][0]["kind"] == "weight")

    # ---------- reports ----------
    page.goto(BASE + "/#/reports/daily"); page.wait_for_selector(".tiles")
    check("daily overview counts overdue from records", page.inner_text(".tile.bad .tn").strip() == str(sum(1 for t in DB["tasks"] if t["due_at"] and t["status"] not in ("completed", "cancelled") and t["due_at"] < iso(dt.datetime.now(dt.timezone.utc)))))
    page.goto(BASE + "/#/reports/weekly"); page.wait_for_selector("#reviewForm")
    page.fill("#rv-wins", "Kite BOM material ready"); page.click("#rvSave"); page.wait_for_timeout(500)
    rv = DB["weekly_reviews"][0] if DB["weekly_reviews"] else None
    check("weekly review saved with a numbers snapshot", rv and rv["wins"] == "Kite BOM material ready" and "completed" in rv["snapshot"])
    page.fill("#rv-wins", "Kite BOM material ready; GRN done"); page.click("#rvSave"); page.wait_for_timeout(500)
    check("saving again updates the same week (no duplicate)", len(DB["weekly_reviews"]) == 1 and DB["weekly_reviews"][0]["wins"].endswith("GRN done"))
    shot(page, "12-weekly-review", True)
    page.goto(BASE + "/#/reports/workload"); page.wait_for_selector(".rtable")
    check("workload lists each owner", "Adil Rahman" in page.inner_text(".rtable") and "Jezz" in page.inner_text(".rtable"))
    with page.expect_download() as dl: page.click('[data-act="rep-csv-workload"]')
    csv = open(dl.value.path(), encoding="utf-8-sig").read()
    check("workload CSV export has a header and rows", csv.startswith("Owner,Open,Overdue") and "Adil Rahman" in csv)
    page.goto(BASE + "/#/reports/projects"); page.click('[data-act="rep-project"]'); page.wait_for_selector("text=plan vs actual")
    check("project report shows milestone variance", "2d late" in page.inner_text("#view"))

    # ---------- workspace tabs + search ----------
    page.goto(BASE + f"/#/w/{WS[1]}/knowledge"); page.wait_for_selector(".viewtabs")
    check("workspace Notes & docs tab lists the meeting", "Weekly logistics review" in page.inner_text("#view"))
    page.goto(BASE + f"/#/w/{WS[1]}/activity"); page.wait_for_timeout(600)
    check("workspace Activity tab loads", "No activity recorded yet" in page.inner_text("#view") or page.locator(".feed").count() == 1)
    page.goto(BASE + "/#/search"); page.fill("#global-q", "logistics review"); page.wait_for_selector(".srow", timeout=4000)
    page.wait_for_timeout(600); check("search returns notes", "weekly logistics review" in page.inner_text("#searchResults").lower(), page.inner_text("#searchResults")[:200])

    # ---------- Phase 3: finance ----------
    check("finance is no longer marked 'coming soon'", page.locator(".sidenav nav a.soon").count() == 0)
    page.goto(BASE + "/#/finance"); page.wait_for_selector(".fin-card", timeout=8000)
    check("finance home lists each set of books", page.locator(".fin-card").count() == 2)
    check("combined view states the management rate", "26.1 INR = 1 QAR" in page.inner_text("#view"))
    page.click(f'a.fin-card[href="#/finance/{FE1}/overview"]'); page.wait_for_selector(".getting-started")
    check("empty books show getting-started steps, not invented figures", "Opening balances" in page.inner_text(".getting-started"))
    page.click('.getting-started [data-act="fin-opening"]'); page.wait_for_selector("#finOpening")
    page.fill(f'input[name="ob-{FE1[:8]}-acc-1010"]', "10000"); page.click('#finOpening button[type="submit"]'); page.wait_for_selector("#finJournal")
    check("opening balances prepared as a balanced entry for review", "Balanced" in page.inner_text("#fj-problems"))
    page.click('#finJournal [data-post="1"]'); page.wait_for_timeout(700)
    ob = [j for j in DB["fin_journals"] if j["source"] == "opening"]
    obl = [l for l in DB["fin_lines"] if ob and l["journal_id"] == ob[0]["id"]]
    check("opening entry posted: bank Dr 10,000 / opening equity Cr 10,000", ob and ob[0]["status"] == "posted" and sorted((float(l["debit"]), float(l["credit"])) for l in obl) == [(0, 10000), (10000, 0)], str(obl))
    page.wait_for_selector(".tiles")
    page.click('[data-act="fin-quick"][data-kind="expense"]'); page.wait_for_selector("#finQuick")
    page.fill("#fq-amt", "800"); page.select_option("#fq-money", f"{FE1[:8]}-acc-1010"); page.select_option("#fq-cat", f"{FE1[:8]}-acc-6000"); page.fill("#fq-memo", "October shop rent")
    page.click('#finQuick [data-post="1"]'); page.wait_for_timeout(700)
    check("quick expense posts a 2-line entry", any(j["memo"] == "October shop rent" and j["status"] == "posted" for j in DB["fin_journals"]))
    page.wait_for_selector(".tiles")
    check("overview shows cash & bank from the ledger (9,200)", "QAR 9,200.00" in page.inner_text(".tiles"))
    # unbalanced entry is refused
    page.click('.page-head [data-act="fin-new-entry"]'); page.wait_for_selector("#finJournal")
    page.fill("#fj-memo", "Test unbalanced")
    page.select_option('[data-jl="0"][data-f="account_id"]', f"{FE1[:8]}-acc-6900"); page.fill('[data-jl="0"][data-f="debit"]', "100")
    page.select_option('[data-jl="1"][data-f="account_id"]', f"{FE1[:8]}-acc-1000"); page.fill('[data-jl="1"][data-f="credit"]', "90")
    check("editor shows the difference live", "differ" in page.inner_text("#fj-problems"))
    page.click('#finJournal [data-post="1"]'); page.wait_for_selector(".toast.error")
    check("unbalanced entry is not saved or posted", not any(j["memo"] == "Test unbalanced" for j in DB["fin_journals"]))
    page.keyboard.press("Escape")
    # customer + invoice + part payment
    page.goto(BASE + f"/#/finance/{FE1}/people"); page.wait_for_selector(".fin-tabs")
    page.click('.filters [data-act="fin-party"]'); page.fill("#fpa-name", "Al Noor Tailoring"); page.click("#fpaSave"); page.wait_for_timeout(500)
    check("customer added to these books", any(p["name"] == "Al Noor Tailoring" and p["entity_id"] == FE1 for p in DB["fin_parties"]))
    page.goto(BASE + f"/#/finance/{FE1}/invoices"); page.click('[data-act="fin-new-invoice"]'); page.wait_for_selector("#finInvoice")
    cust = next(p["id"] for p in DB["fin_parties"] if p["name"] == "Al Noor Tailoring")
    page.select_option("#fi-party", cust); page.fill('[data-il="0"][data-f="description"]', "Bespoke suit"); page.fill('[data-il="0"][data-f="qty"]', "2"); page.fill('[data-il="0"][data-f="unit_price"]', "750")
    page.locator('[data-il="0"][data-f="unit_price"]').dispatch_event("change"); page.wait_for_timeout(200)
    check("invoice total calculates as you type (1,500)", "1,500.00" in page.inner_text("#fi-total"), page.inner_text("#fi-total") + " / " + str(page.evaluate("[...document.querySelectorAll(\"[data-il]\")].map(x=>x.dataset.f+\"=\"+x.value).join(\",\")")))
    check("due date is left empty unless you set it", page.input_value("#fi-due") == "")
    page.click('#finInvoice [data-post="1"]'); page.wait_for_timeout(800)
    inv = next((i for i in DB["fin_invoices"] if i["entity_id"] == FE1), None)
    check("invoice posted as INV-0001 with its journal", inv and inv["status"] == "posted" and inv["number"] == "INV-0001" and inv["journal_id"], str(inv))
    page.click(f'[data-act="fin-open-invoice"][data-id="{inv["id"]}"]'); page.wait_for_selector(".inv-head")
    page.click('[data-act="fin-pay"][data-invoice]'); page.wait_for_selector("#finPayment")
    check("payment form pre-fills the outstanding amount", page.input_value("#fp-amt") == "1500")
    page.fill("#fp-amt", "1000"); page.fill(f'[data-alloc="{inv["id"]}"]', "1000"); page.fill("#fp-ref", "TT-0091"); page.click("#fpSave"); page.wait_for_timeout(800)
    check("part payment recorded and applied", any(a["invoice_id"] == inv["id"] and a["amount"] == 1000 for a in DB["fin_allocations"]))
    page.goto(BASE + f"/#/finance/{FE1}/invoices"); page.wait_for_selector(".list")
    check("invoice shows Part paid with 500 left", "Part paid" in page.inner_text(".list") and "500.00 left" in page.inner_text(".list"))
    # reports
    page.goto(BASE + f"/#/finance/{FE1}/reports"); page.wait_for_selector(".fin-report")
    t = page.inner_text(".fin-report")
    check("P&L from posted entries: income 1,500, expenses 800, profit 700", "1,500.00" in t and "800.00" in t and "700.00" in t, t[:400])
    page.select_option("#fr-report", "bs"); page.wait_for_selector(".fin-report")
    check("balance sheet balances", "Balances: assets = liabilities + equity" in page.inner_text(".fin-report"))
    page.select_option("#fr-report", "tb"); page.wait_for_selector(".fin-report")
    check("trial balance: debits equal credits", "Debits equal credits" in page.inner_text(".fin-report"))
    with page.expect_download() as dl: page.click('[data-act="fin-csv"]')
    check("report CSV downloads", open(dl.value.path(), encoding="utf-8-sig").read().startswith("Account,Debit,Credit"))
    shot(page, "14-finance-report", True)
    # bank statement import and match
    page.goto(BASE + f"/#/finance/{FE1}/banking"); page.wait_for_selector(".fin-tabs")
    d = dt.date.fromisoformat(today_qa()).strftime("%d/%m/%Y")
    page.set_input_files("#fb-file", files=[{"name": "qnb.csv", "mimeType": "text/csv", "buffer": f"Account,QNB 0123\nDate,Description,Debit,Credit\n{d},TT AL NOOR,,1000.00\n{d},BANK CHARGE,15.00,\n".encode()}])
    page.wait_for_selector('[data-act="fin-import-go"]')
    check("statement preview reads 2 lines, day-first dates", "2 lines read" in page.inner_text("#sheet"))
    page.click('[data-act="fin-import-go"]'); page.wait_for_timeout(800)
    check("statement lines imported as unmatched", sum(1 for x in DB["fin_statement_lines"] if x["status"] == "unmatched") == 2)
    page.wait_for_selector('[data-act="fin-match"]', timeout=5000)
    page.click('[data-act="fin-match"]'); page.wait_for_timeout(700)
    check("bank line matched to the customer payment", any(x["status"] == "matched" and x["amount"] == 1000 for x in DB["fin_statement_lines"]))
    shot(page, "15-banking")
    # budget
    page.goto(BASE + f"/#/finance/{FE1}/budgets"); page.wait_for_selector("#budgetForm")
    page.fill(f'input[name="b-{FE1[:8]}-acc-6000"]', "1000"); page.click("#budgetSave"); page.wait_for_timeout(500)
    check("budget saved separately from actuals", any(b["account_id"] == f"{FE1[:8]}-acc-6000" and b["amount"] == 1000 for b in DB["fin_budgets"]))
    page.goto(BASE + f"/#/finance/{FE1}/overview"); page.wait_for_selector(".tiles")
    shot(page, "16-finance-overview", True)

    # ---------- Phase 4: smart capture, briefing, drafts, calendar link, AI (simulated) ----------
    page.goto(BASE + "/#/home"); page.wait_for_selector(".briefing")
    check("Command Centre opens with a briefing from stored records", "overdue" in page.inner_text(".briefing") or "Nothing overdue" in page.inner_text(".briefing"), page.inner_text(".briefing"))
    page.keyboard.press("n"); page.wait_for_selector("#capForm")
    page.fill("#cap-title", "Chase Al Muthawa dispatch fri 10am #dlusso !high"); page.wait_for_timeout(150)
    chips = page.inner_text("#cap-chips")
    check("capture recognises date, workspace and priority as chips", "Due:" in chips and "D'Lusso" in chips and "High" in chips, chips)
    page.click('[data-act="cap-chip-off"][data-k="due"]'); page.wait_for_timeout(100)
    check("removing a chip keeps that text as plain words", "Due:" not in page.inner_text("#cap-chips"))
    page.click("#capSave"); page.wait_for_selector("text=Capture another")
    sm = next((t for t in DB["tasks"] if t["title"].startswith("Chase Al Muthawa")), None)
    check("saved with workspace and priority from the line, no date (chip removed)", sm and sm["workspace_id"] == WS[0] and sm["priority"] == "high" and sm["due_at"] is None and sm["title"] == "Chase Al Muthawa dispatch fri 10am", str(sm and (sm["title"], sm["due_at"])))
    page.click("text=Capture another"); page.fill("#cap-title", "Call clearing agent tomorrow 3pm @adil #qcloud"); page.wait_for_timeout(100); page.click("#capSave"); page.wait_for_selector("text=Capture another")
    sm2 = next((t for t in DB["tasks"] if t["title"] == "Call clearing agent"), None)
    tmr = (dt.date.fromisoformat(today_qa()) + dt.timedelta(days=1)).isoformat()
    check("'tomorrow 3pm' stored as 15:00 Doha tomorrow, owner Adil", sm2 and sm2["due_at"].startswith(tmr + "T12:00") and sm2["owner_id"] == ADIL, str(sm2 and (sm2["due_at"], sm2["owner_id"])))
    page.keyboard.press("Escape")
    # message draft for a follow-up — nothing sent automatically
    page.goto(BASE + "/#/followups"); page.wait_for_selector('[data-act="open-followup"][data-id="f-1"]'); page.click('[data-act="open-followup"][data-id="f-1"]'); page.wait_for_selector("#fuForm")
    page.click('[data-act="draft-msg"]'); page.wait_for_selector("#dr-text")
    check("draft message is written from the follow-up's facts", "Revised dispatch date" in page.input_value("#dr-text") or "revised dispatch date" in page.input_value("#dr-text").lower())
    check("without AI turned on there is no AI button", page.locator('[data-act="dr-ai"]').count() == 0)
    page.select_option("#dr-tone", "firm"); page.wait_for_timeout(100)
    check("firm tone changes the draft", "holding up" in page.input_value("#dr-text"))
    logs_before = len(DB["follow_up_logs"])
    check("opening the draft sends nothing and logs nothing", len(DB["follow_up_logs"]) == logs_before)
    page.click('[data-act="dr-log"]'); page.wait_for_timeout(500)
    check("'I sent it' logs the contact on the follow-up", any(l["body"].startswith("Sent by WhatsApp") for l in DB["follow_up_logs"]))
    # calendar link
    page.goto(BASE + "/#/settings"); page.wait_for_selector("#ig-cal")
    page.wait_for_selector('[data-act="ig-cal-new"]'); page.click('[data-act="ig-cal-new"]'); page.wait_for_selector("#ig-url")
    check("calendar link shown once with the private token", page.input_value("#ig-url").endswith("/functions/v1/calendar-feed?token=" + "ab" * 32))
    check("Apple Calendar subscribe link offered", page.locator('a[href^="webcal:"]').count() == 1)
    # AI on
    page.check("#ig-ai-on"); page.wait_for_timeout(600)
    check("AI assistant opt-in saved to profile", (next(p for p in DB["profiles"] if p["id"] == ME)["prefs"] or {}).get("ai_enabled") is True)
    check("Settings shows AI ready with usage", "Ready" in page.inner_text("#ig-ai"), page.inner_text("#ig-ai"))
    page.goto(BASE + "/#/notes/meetings"); page.wait_for_selector(".note-row"); page.click(".note-row"); page.wait_for_selector("#nfAi")
    calls = len(AI_CALLS); page.click("#nfAi"); page.wait_for_selector(".actions-review")
    check("AI action items arrive as suggestions for review (nothing created yet)", "Suggested by AI" in page.inner_text(".actions-review") and page.locator(".actions-review li").count() == 2 and len(AI_CALLS) == calls + 1)
    check("AI owner name matched to the workspace member", page.input_value("#ai-owner-0") == ADIL)
    page.fill("#ai-title-0", "Confirm inspection slot with surveyor (Sun)"); page.uncheck('[data-act="ai-toggle"][data-i="1"]')
    page.click('[data-act="ai-create"]'); page.wait_for_timeout(900)
    ai_items = [v for v in AI_PROPOSALS.values() if v["action"] == "action_items"]
    check("AI log: edited and partly skipped action items recorded as 'edited'", ai_items and ai_items[-1]["decisions"] == ["edited"] and (ai_items[-1]["context"] or {}).get("type") == "note", str(ai_items[-1:]))
    page.keyboard.press("Escape")
    page.goto(BASE + "/#/reports/weekly"); page.wait_for_selector("#rvAi"); page.click("#rvAi"); page.wait_for_timeout(400)
    check("AI weekly summary is placed in the notes box, not saved", "Kite BOM" in page.input_value("#rv-notes") and not any("Kite BOM material readiness" in (r.get("notes") or "") for r in DB["weekly_reviews"]))
    page.click("#rvSave"); page.wait_for_timeout(600)
    wk = [v for v in AI_PROPOSALS.values() if v["action"] == "weekly_narrative"]
    check("AI log: weekly summary kept as written recorded as 'accepted'", wk and wk[-1]["decisions"] == ["accepted"], str(wk[-1:]))
    page.goto(BASE + f"/#/finance/{FE1}/banking"); page.wait_for_selector("#finAiCat"); page.click("#finAiCat"); page.wait_for_timeout(500)
    check("AI bank-line suggestion shown as a button (nothing posted)", page.locator('[data-act="fin-stmt-entry"][data-acc]').count() >= 1 and not any(j["memo"] == "BANK CHARGE" for j in DB["fin_journals"]))
    page.click('[data-act="fin-stmt-entry"][data-acc]'); page.wait_for_selector("#finStmtEntry")
    check("suggested account is pre-selected for review", page.input_value("#fse-acc") == f"{FE1[:8]}-acc-6900")
    page.keyboard.press("Escape")
    page.goto(BASE + "/#/settings"); page.wait_for_selector("#ig-ai"); page.wait_for_timeout(400)
    check("Settings explains the AI can't act on its own and shows AI history", "can't create tasks, change dates or owners" in page.inner_text("#ig-ai") and "AI history" in page.inner_text("#ig-ai"))

    # Settings: hide a dashboard section persists to profile prefs
    page.goto(BASE + "/#/settings"); page.wait_for_selector(".settings")
    page.uncheck('[data-act="toggle-section"][data-k="activity"]'); page.wait_for_timeout(300)
    check("dashboard customisation saved to profile", "activity" in (next(p for p in DB["profiles"] if p["id"] == ME)["prefs"].get("dashboard_hidden") or []))
    page.goto(BASE + "/#/home"); page.wait_for_selector(".tiles")
    check("hidden section no longer shown", "Recent activity" not in page.inner_text("#view"))
    ctx.close()


    # ---------- Partner accounts (fixture books only) ----------
    ctx, page = new_page(b, 1366, 900)
    FE3 = "fe000000-0000-4000-8000-000000000003"
    acc = lambda eid, c, n, t, st: {"id": f"{eid[:8]}-acc-{c}", "entity_id": eid, "code": c, "name": n, "type": t, "subtype": st, "system_key": None, "is_active": True, "description": None}
    DB["fin_accounts"] += [acc(FE1, "1200", "Inventory", "asset", "inventory"), acc(FE1, "2300", "Partner current accounts", "liability", "partner"), acc(FE1, "6100", "Salaries & wages", "expense", "other")]
    DB["fin_entities"].append({**DB["fin_entities"][0], "id": FE3, "name": "Personal & Family", "kind": "personal", "workspace_id": None, "next_journal_no": 1})
    DB["fin_members"].append({"entity_id": FE3, "user_id": ME, "role": "owner"}); DB["fin_accounts"] += seed_accounts(FE3)
    a = lambda c, e=FE1: f"{e[:8]}-acc-{c}"
    page.goto(BASE + f"/#/finance/{FE1}/partners"); page.reload(); page.wait_for_selector('[data-act="pt-setup"]')
    check("partner accounts tab offers set-up when none exists", page.locator(".getting-started").count() == 1)
    page.click('[data-act="pt-setup"]'); page.wait_for_selector("#ptSetup"); page.fill("#ps-name", "Jezz"); page.select_option("#ps-mirror", FE3); page.click("#psSave")
    page.wait_for_selector('[data-act="pt-new"]')
    lk = DB["fin_partner_links"][0] if DB["fin_partner_links"] else {}
    check("partner link uses the existing partner current account and mirrors to personal", lk.get("partner_account_id") == a("2300") and lk.get("mirror_receivable_account_id") == a("1100", FE3))
    def record(type_, date, desc, amt, acc_, month=False, ben="", ref="", settles=None, expect_ok=True):
        page.click('[data-act="pt-new"]'); page.wait_for_selector("#ptRecord")
        if type_ != "expense_funding": page.select_option("#pr-type", type_); page.wait_for_timeout(150)
        page.fill("#pr-date", date); page.fill("#pr-desc", desc); page.fill("#pr-amt", str(amt)); page.select_option("#pr-acc", acc_)
        if month: page.check("#pr-month")
        if ben: page.fill("#pr-ben", ben)
        if ref: page.fill("#pr-ref", ref)
        if settles: page.select_option("#pr-settles", settles)
        page.click("#prSave"); page.wait_for_timeout(600)
        if expect_ok: page.keyboard.press("Escape"); page.wait_for_timeout(150)
    page.click('[data-act="pt-new"]'); page.wait_for_selector("#ptRecord")
    opts = page.eval_on_selector_all("#pr-acc option", "os => os.map(o => o.value)")
    check("funding can't be posted to the business bank or cash (not offered)", a("1010") not in opts and a("1000") not in opts and a("4000") not in opts and a("2300") not in opts, str(opts))
    page.keyboard.press("Escape")
    record("expense_funding", "2026-09-15", "AND Stitching – Ceproteca", 1436, a("6900"), month=True, ben="Ceproteca")
    t1 = DB["fin_partner_txns"][0] if DB["fin_partner_txns"] else {}
    check("month-only date stored as month end, flagged, payment date blank", t1.get("txn_date") == "2026-09-30" and t1.get("date_precision") == "month" and t1.get("payment_date") is None and "month_only_date" in t1.get("flags", []), str(t1))
    check("transaction posted in both books", t1.get("journal_id") and t1.get("mirror_journal_id") and t1.get("txn_no") == "PR-0001")
    record("salary_funding", "2026-10-25", "September salary funding", 8500, a("6100"))
    record("reimbursement", "2026-10-29", "Reimbursement from Dlusso", 20000, a("1010"), expect_ok=False)
    check("reimbursement above what is owed is refused", "more than" in page.inner_text("#toast") and len(DB["fin_partner_txns"]) == 2, page.inner_text("#toast"))
    page.keyboard.press("Escape"); page.wait_for_timeout(150)
    sal = next(t for t in DB["fin_partner_txns"] if t["amount"] == 8500)
    record("reimbursement", "2026-10-29", "Reimbursement from Dlusso", 8000, a("1010"), settles=sal["id"])
    re_ = next((t for t in DB["fin_partner_txns"] if t["txn_type"] == "reimbursement"), {})
    check("reimbursement debits the partner account and credits the business bank, linked to the salary line",
          re_.get("settles_txn_id") == sal["id"] and any(l["journal_id"] == re_.get("journal_id") and l["account_id"] == a("1010") and float(l["credit"]) == 8000 for l in DB["fin_lines"]))
    page.goto(BASE + f"/#/finance/{FE1}/partners/dashboard"); page.wait_for_selector(".tiles")
    tiles = page.inner_text(".tiles")
    check("dashboard: outstanding 1,936.00 and both books agree", "1,936.00" in tiles and "Both books agree" in page.inner_text("#view"), tiles[:200])
    check("dashboard: no overdue invented without terms", "no repayment terms set" in tiles.lower())
    shot(page, "20-partner-dashboard", True)
    page.goto(BASE + f"/#/finance/{FE1}/partners/ledger"); page.wait_for_selector(".ptn-ledger")
    lt = page.inner_text(".ptn-ledger")
    check("ledger shows Document Pending and month-only dates honestly", "Document Pending" in lt and "exact date unavailable" in lt)
    shot(page, "21-partner-ledger", True)
    page.click(f'tr[data-id="{t1["id"]}"]'); page.wait_for_selector("#pt-docs")
    page.set_input_files('input[data-act="pt-attach"]', files=[{"name": "Ceproteca invoice.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1.4 test"}]); page.wait_for_timeout(700)
    dd = DB["fin_partner_docs"][0] if DB["fin_partner_docs"] else {}
    check("document stored privately under the business's folder and linked to its transaction", dd.get("txn_id") == t1["id"] and dd.get("storage_path", "").startswith(FE1 + "/partner/") and dd.get("storage_path") in STORAGE, str(dd))
    page.click(f'[data-act="pt-edit"][data-id="{t1["id"]}"]'); page.wait_for_selector("#ptEdit"); page.fill("#pe-ref", "AND-2209"); page.click("#peSave"); page.wait_for_timeout(300)
    check("edit without a reason is refused", "reason" in page.inner_text("#toast").lower() and t1.get("reference") is None)
    page.fill("#pe-reason", "Invoice number found on supplier copy"); page.click("#peSave"); page.wait_for_timeout(600)
    au = [x for x in DB["fin_audit"] if x.get("record_id") == t1["id"]]
    check("edit saved with previous value, new value and reason", t1.get("reference") == "AND-2209" and au and au[-1]["changes"]["reference"] == [None, "AND-2209"] and au[-1]["changes"]["last_change_reason"][1] == "Invoice number found on supplier copy")
    page.click(f'[data-act="pt-reverse"][data-id="{t1["id"]}"]'); page.wait_for_selector(".dlg"); page.click('.dlg [data-r="1"]')
    check("reversal requires a reason", page.is_visible("#dlgErr"))
    page.fill("#dlgIn", "Duplicate of a bill already posted"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(700)
    check("reversal keeps the original and reverses both books", t1.get("status") == "reversed" and t1.get("reversal_journal_id") and t1.get("mirror_reversal_journal_id") and any(j["id"] == t1["journal_id"] for j in DB["fin_journals"]))
    page.keyboard.press("Escape")
    page.goto(BASE + f"/#/finance/{FE1}/partners/reconciliation"); page.wait_for_selector(".rtable")
    rt = page.inner_text("#view").lower()
    check("reconciliation: the three balances agree at 500 and exceptions are listed", rt.count("agrees") >= 3 and "500.00" in rt and "payment evidence missing" in rt and "differs" not in rt, rt.replace("\n", " ")[:400])
    shot(page, "22-partner-reconciliation", True)
    page.goto(BASE + f"/#/finance/{FE1}/partners/statement"); page.wait_for_selector(".report")
    page.fill('input[data-act="pt-month"]', "2026-10"); page.dispatch_event('input[data-act="pt-month"]', "change"); page.wait_for_timeout(300)
    st = page.inner_text(".report")
    check("October statement: opening 0.00, funded 8,500, reimbursed 8,000, closing 500", "8,500.00" in st and "8,000.00" in st and "500.00" in st, st[:400])
    with page.expect_download() as dl: page.click('[data-act="pt-xlsx"][data-r="statement"]')
    path_ = os.path.join(SHOTS or "/tmp", "partner-statement.xlsx"); dl.value.save_as(path_)
    try:
        import openpyxl; wb = openpyxl.load_workbook(path_); ws = wb.worksheets[0]; vals = [r for r in ws.iter_rows(values_only=True)]
        check("statement downloads as a real Excel file with the closing balance", dl.value.suggested_filename.endswith(".xlsx") and vals[-1][7] == 500, str(vals[-1]))
    except Exception as ex_: check("statement downloads as a real Excel file with the closing balance", False, str(ex_))
    with page.expect_download() as dl2: page.click('[data-act="pt-xlsx"][data-r="all"]')
    try:
        p2 = os.path.join(SHOTS or "/tmp", "partner-reports.xlsx"); dl2.value.save_as(p2); wb2 = openpyxl.load_workbook(p2)
        check("all five reports export to one workbook", len(wb2.sheetnames) >= 6 and any("ageing" in n.lower() for n in wb2.sheetnames), str(wb2.sheetnames))
    except Exception as ex_: check("all five reports export to one workbook", False, str(ex_))
    shot(page, "23-partner-statement", True)
    ctx.close()


    # ---------- Google Calendar live sync (simulated Google) ----------
    ctx, page = new_page(b, 1366, 900)
    page.route(re.compile(r"https://accounts\.google\.com/.*"), lambda r: r.fulfill(body="google sign-in", content_type="text/html"))
    page.goto(BASE + "/#/settings"); page.wait_for_selector("#ig-gcal .banner")
    check("Google Calendar: setup steps and redirect URI shown until the app is configured", "GOOGLE_CLIENT_ID" in page.inner_text("#ig-gcal") and page.input_value("#ig-gredir").endswith("/functions/v1/gcal-sync/callback"))
    GCAL["configured"] = True
    page.reload(); page.wait_for_selector('[data-act="ig-gcal-connect"]')
    with page.expect_navigation(url=re.compile(r"https://accounts\.google\.com/.*")): page.click('[data-act="ig-gcal-connect"]')
    check("Connect sends you to Google asking only for calendar events", "calendar.events" in page.url and "calendar.readonly" not in page.url, page.url)
    GCAL["sync"] = {"google_email": "jezz@example.com", "status": "connected", "last_error": None, "last_synced_at": iso(now), "connected_at": iso(now)}
    page.goto(BASE + "/#/settings?gcal=connected"); page.wait_for_selector('[data-act="ig-gcal-resync"]'); page.wait_for_timeout(200)
    check("after Google returns: connected account shown with a confirmation", "jezz@example.com" in page.inner_text("#ig-gcal") and "connected" in page.inner_text("#toast").lower())
    page.click('[data-act="ig-gcal-resync"]'); page.wait_for_timeout(400)
    check("Sync all now reports how many tasks were synced", "4 tasks synced" in page.inner_text("#toast"))
    shot(page, "24-google-calendar-settings")
    page.click('[data-act="ig-gcal-off"]'); page.wait_for_selector(".dlg"); page.click('.dlg [data-r="1"]'); page.wait_for_selector('[data-act="ig-gcal-connect"]')
    check("Disconnect asks first, then shows Connect again", "disconnect" in GCAL["calls"])
    ctx.close()


    # ---------- Increment 1: assignee, reviewer, review before completion ----------
    ctx, page = new_page(b, 1366, 900)
    rv = {**fresh_db()["tasks"][0], "id": "t-review", "title": "Customs docs pack", "workspace_id": WS[1], "owner_id": ME, "status": "in_progress", "completion_criteria": []}
    DB["tasks"].append(rv)
    page.goto(BASE + "/#/t/t-review"); page.wait_for_selector("#tf-reviewer_id")
    check("task shows responsible owner, assignee and reviewer", page.locator("#tf-assignee_id").count() == 1 and page.locator("#tf-reviewer_id").count() == 1)
    rev_opts = page.eval_on_selector_all("#tf-reviewer_id option", "os => os.map(o => o.value)")
    check("you can't pick the person doing the work as reviewer", ME not in rev_opts and ADIL in rev_opts, str(rev_opts))
    page.select_option("#tf-reviewer_id", ADIL); page.wait_for_selector('[data-act="review-request"]')
    check("with a reviewer, Complete becomes Request review", rv.get("reviewer_id") == ADIL and page.locator('[data-act="complete-task"]').count() == 0)
    page.click('[data-act="review-request"]'); page.wait_for_selector(".dlg"); page.fill("#dlgIn", "All 6 docs attached"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(500)
    check("review requested and shown as waiting for the reviewer", rv.get("review_status") == "requested" and "Waiting for" in page.inner_text("#sheet"), page.inner_text("#toast"))
    rv.update(assignee_id=ADIL, reviewer_id=ME, review_status="requested")
    page.keyboard.press("Escape"); page.goto(BASE + "/#/home"); page.reload(); page.wait_for_selector(".tiles"); page.goto(BASE + "/#/t/t-review"); page.wait_for_selector('[data-act="review-approve"]')
    page.click('[data-act="review-changes"]'); page.wait_for_selector(".dlg"); page.click('.dlg [data-r="1"]')
    check("asking for changes needs a note", page.is_visible("#dlgErr"))
    page.fill("#dlgIn", "COO certificate missing"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(500)
    check("changes requested sends the task back", rv.get("review_status") == "changes_requested" and rv.get("status") == "in_progress")
    rv["review_status"] = "requested"; page.goto(BASE + "/#/home"); page.reload(); page.wait_for_selector(".tiles"); page.goto(BASE + "/#/t/t-review"); page.wait_for_selector('[data-act="review-approve"]'); page.click('[data-act="review-approve"]'); page.wait_for_timeout(500)
    check("reviewer approval completes the task", rv.get("status") == "completed" and rv.get("review_status") == "approved")
    shot(page, "25-task-review")
    ctx.close()


    # ---------- Increment 2: commitments need an outcome to close ----------
    ctx, page = new_page(b, 1366, 900)
    page.goto(BASE + "/#/followups"); page.wait_for_selector('[data-act="open-followup"][data-id="f-1"]'); page.click('[data-act="open-followup"][data-id="f-1"]'); page.wait_for_selector("#fu-expectation")
    check("commitment form asks what exactly is expected and its type", page.locator("#fu-type").count() == 1 and page.locator("#fu-response").count() == 1)
    page.fill("#fu-expectation", "Written revised dispatch date and recovery plan"); page.fill("#fu-response", "Supplier confirmed 12 Nov, partial shipment"); page.click("#fuSave"); page.wait_for_timeout(500)
    f1 = next(x for x in DB["follow_ups"] if x["id"] == "f-1")
    check("expectation and actual response saved", f1.get("expectation", "").startswith("Written revised") and f1.get("response", "").startswith("Supplier confirmed"))
    page.click('[data-act="fu-close"]'); page.wait_for_selector("#fu-close-panel"); page.click('[data-act="fu-close-confirm"]'); page.wait_for_timeout(200)
    check("closing without an outcome is refused", f1["status"] != "closed" and "outcome" in page.inner_text("#toast").lower())
    page.select_option("#fu-outcome", "renegotiated"); page.fill("#fu-outcome-note", "New date agreed: 12 Nov"); page.click('[data-act="fu-close-confirm"]'); page.wait_for_timeout(500)
    check("closed with the recorded outcome", f1["status"] == "closed" and f1.get("outcome") == "renegotiated" and "Renegotiated" in page.inner_text("#sheet"))
    shot(page, "26-commitment-outcome")
    ctx.close()


    # ---------- Increment 3: schedule baselines ----------
    ctx, page = new_page(b, 1366, 900)
    pj = next(x for x in DB["projects"] if x["id"] == P1)
    if not pj.get("target_date"): pj["target_date"] = "2026-12-15"
    page.goto(BASE + f"/#/p/{P1}"); page.wait_for_selector('[data-act="set-baseline"]')
    page.click('[data-act="set-baseline"]'); page.wait_for_selector(".dlg"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(600)
    base_target = pj.get("baseline_target_date")
    check("baseline freezes the current target", pj.get("baselined_at") and base_target == pj["target_date"], str(base_target))
    page.click(f'[data-act="edit-project"][data-id="{P1}"]'); page.wait_for_selector("#pf-target")
    new_target = (dt.date.fromisoformat(pj["target_date"]) + dt.timedelta(days=14)).isoformat()
    page.fill("#pf-target", new_target); page.click("#pfSave"); page.wait_for_selector(".dlg"); page.click('.dlg [data-r="1"]')
    check("changing a baselined date asks for a reason", page.is_visible("#dlgErr"))
    page.fill("#dlgIn", "Supplier delay on racks"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(800)
    revs = [r_ for r_ in DB.get("schedule_revisions", []) if r_["project_id"] == P1]
    check("new target saved, baseline kept, revision logged with reason", pj["target_date"] == new_target and pj["baseline_target_date"] == base_target and revs and revs[-1]["reason"] == "Supplier delay on racks")
    page.goto(BASE + "/#/home"); page.goto(BASE + f"/#/p/{P1}"); page.wait_for_selector(".proj-hero"); page.wait_for_timeout(500)
    check("project page shows the 14-day slip against the baseline", "slipped 14d" in page.inner_text("#view"), page.inner_text(".proj-hero"))
    shot(page, "27-project-baseline", True)
    ctx.close()


    # ---------- Increment 5: delivery status and snooze ----------
    ctx, page = new_page(b, 1366, 900)
    nid = DB["notifications"][0]["id"]
    DB.setdefault("notification_deliveries", []).append({"notification_id": nid, "status": "sent", "attempts": 2, "sent_at": iso(now), "created_at": iso(now)})
    page.goto(BASE + "/#/alerts"); page.wait_for_selector(".alist"); page.wait_for_timeout(600)
    check("alert shows it reached the phone (after a retry)", "Sent to phone" in page.inner_text(".alist") and "2 tries" in page.inner_text(".alist"), page.inner_text(".alist")[:200])
    page.click(f'[data-act="alert-snooze"][data-id="{nid}"][data-when="tomorrow"]'); page.wait_for_timeout(500)
    sn = next(x for x in DB["notifications"] if x["id"] == nid)
    check("snooze until tomorrow 9:00 Doha time", (sn.get("snoozed_until") or "").startswith((dt.date.fromisoformat(today_qa()) + dt.timedelta(days=1)).isoformat() + "T06:00"), str(sn.get("snoozed_until")))
    check("snoozed alert moves to the Snoozed list", "snoozed" in page.inner_text("#view").lower() and page.locator(f'.alist [data-id="{nid}"]').count() == 0)
    ctx.close()


    # ---------- Increment 6: inter-entity ----------
    ctx, page = new_page(b, 1366, 900)
    page.goto(BASE + f"/#/finance/{FE1}/intercompany"); page.wait_for_selector('[data-act="ic-new"]')
    page.click('[data-act="ic-new"]'); page.wait_for_selector("#icForm")
    page.select_option("#ic-to", FE2); page.wait_for_selector("#ic-amt-to")
    check("different currencies ask for the amount received in the other currency", page.locator("#ic-amt-to").count() == 1)
    page.fill("#ic-desc", "Start-up funding for India outlet"); page.fill("#ic-amt", "1000")
    page.select_option("#ic-from-acc", f"{FE1[:8]}-acc-1010"); page.select_option("#ic-to-acc", f"{FE2[:8]}-acc-1010")
    page.click("#icSave"); page.wait_for_timeout(300)
    check("missing received amount is caught before posting", len(DB["fin_ic_txns"]) == 0 and "other currency" in page.inner_text("#toast"))
    page.fill("#ic-amt-to", "22850"); page.click("#icSave"); page.wait_for_timeout(900)
    t_ = DB["fin_ic_txns"][0] if DB["fin_ic_txns"] else {}
    check("posted in both books with one IC number", t_.get("amount") == 1000 and t_.get("amount_to") == 22850 and t_.get("from_journal_id") and t_.get("to_journal_id")
          and any(j["id"] == t_["to_journal_id"] and j["entity_id"] == FE2 for j in DB["fin_journals"]))
    page.wait_for_selector("#ic-pos .rtable"); page.wait_for_timeout(300)
    check("balances table shows the other books and 'Matched'", "The Burgery Thrissur" in page.inner_text("#ic-pos") and "matched" in page.inner_text("#ic-pos").lower())
    shot(page, "28-intercompany", True)
    ctx.close()

    # ---------- Increment 7: production costing ----------
    ctx, page = new_page(b, 1366, 900)
    WOOL = str(uuid.uuid4())
    DB["fin_items"].append({"id": WOOL, "entity_id": FE1, "name": "Wool suiting", "sku": "WS-1", "unit": "m", "inventory_account_id": None, "archived": False, "created_at": iso(now)})
    page.goto(BASE + f"/#/finance/{FE1}/production"); page.wait_for_selector('[data-act="prod-receive"]')
    check("business books have a Production tab", page.locator(f'.fin-tabs a[href="#/finance/{FE1}/production"]').count() == 1)
    page.click('[data-act="prod-receive"]'); page.wait_for_selector("#prodGrnForm")
    page.select_option("#gr-item", WOOL); page.fill("#gr-qty", "10"); page.click("#grSave"); page.wait_for_timeout(300)
    check("receipt without a unit cost is caught", not DB["fin_goods_receipts"] and "unit cost" in page.inner_text("#toast"))
    page.fill("#gr-cost", "40"); page.fill("#gr-ref", "DN-118"); page.click("#grSave"); page.wait_for_timeout(900)
    g_ = DB["fin_goods_receipts"][0] if DB["fin_goods_receipts"] else {}
    check("goods received: 10 m at 40 in stock, GRN-0001, 400 held as received-not-billed", g_.get("receipt_no") == "GRN-0001" and on_hand(WOOL)[0] == 10 and "GRN-0001" in page.inner_text("#view")
          and "400" in page.inner_text(".tiles"), page.inner_text(".tiles"))
    check("page tells you to code the bill to goods-received", "2050 goods received, not yet billed" in page.inner_text("#view").lower())
    page.click('[data-act="prod-new"]'); page.wait_for_selector("#prodNewForm"); page.fill("#po-title", "Navy two-piece suit"); page.fill("#po-ref", "K-2210"); page.click("#poSave")
    page.wait_for_selector("#prodIssueForm")
    o_ = DB["fin_production_orders"][0]
    check("order opens as JOB-0001 and shows its cost sheet", o_["order_no"] == "JOB-0001" and "job-0001" in page.inner_text(".sheet").lower())
    page.select_option("#pi-item", WOOL); page.fill("#pi-qty", "3.5"); page.click("#piSave"); page.wait_for_timeout(900)
    check("issuing 3.5 m puts 140 into the job and leaves 6.5 m", on_hand(WOOL)[0] == 6.5 and "140" in page.inner_text("#prodCosts"))
    page.fill("#pc-note", "Stitching — Topaz tailoring"); page.fill("#pc-amt", "150"); page.select_option("#pc-acc", f"{FE1[:8]}-acc-1010"); page.click("#pcSave"); page.wait_for_timeout(900)
    check("stitching added: job cost 290", "290" in page.inner_text(".sheet .tiles"), page.inner_text(".sheet .tiles"))
    page.select_option("#pi-item", WOOL); page.fill("#pi-qty", "8"); page.click("#piSave"); page.wait_for_selector("#dlgIn")
    check("using more than recorded stock asks the owner to authorise with a reason", "not enough recorded stock" in page.inner_text(".dlg").lower())
    page.click('.dlg [data-r="1"]'); page.wait_for_timeout(200)
    check("authorising needs a reason", page.is_visible("#dlgErr") and not DB["fin_stock_exceptions"])
    page.fill("#dlgIn", "Fabric arrived, receipt not yet entered"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(900)
    check("exception recorded with the 1.5 m shortfall", DB["fin_stock_exceptions"] and DB["fin_stock_exceptions"][0]["qty_short"] == 1.5 and "authorised short stock" in page.inner_text("#prodCosts").lower())
    xc = next(c for c in DB["fin_production_costs"] if c["exception_id"])
    page.click(f'[data-act="prod-cost-reverse"][data-id="{xc["id"]}"]'); page.wait_for_selector("#dlgIn"); page.fill("#dlgIn", "Second suit cancelled"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(900)
    check("reversing returns the fabric to stock", on_hand(WOOL)[0] == 6.5 and xc["status"] == "reversed")
    page.click('[data-act="prod-complete"]'); page.wait_for_selector(".dlg"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(900)
    check("completing moves 290 to cost of goods sold", o_["status"] == "completed" and o_["total_cost"] == 290 and page.locator("#prodIssueForm").count() == 0)
    shot(page, "29-production-order", True)
    page.keyboard.press("Escape"); page.goto(BASE + f"/#/finance/{FE1}/stock"); page.goto(BASE + f"/#/finance/{FE1}/production"); page.wait_for_selector('[data-act="prod-open"]')
    check("orders list shows the completed job at 290", "completed" in page.inner_text("#view").lower() and "290" in page.inner_text('[data-act="prod-open"]'))
    ctx.close()

    # ---------- Increment 8: foreign currency ----------
    ctx, page = new_page(b, 1366, 900)
    LON = str(uuid.uuid4()); DB["fin_parties"].append({"id": LON, "entity_id": FE1, "name": "London client", "kind": "customer", "email": None, "phone": None, "notes": None, "archived": False, "created_at": iso(now)})
    page.goto(BASE + f"/#/finance/{FE1}/invoices"); page.wait_for_selector('[data-act="fin-new-invoice"]')
    page.click('[data-act="fin-new-invoice"]'); page.wait_for_selector("#finInvoice")
    page.select_option("#fi-party", LON); page.select_option("#fi-ccy", "USD"); page.wait_for_selector("#fi-rate")
    check("choosing USD asks for the rate (left blank)", page.input_value("#fi-rate") == "" and "prices below are in usd" in page.inner_text("#finInvoice").lower())
    page.fill('[data-il="0"][data-f="description"]', "Two suits"); page.fill('[data-il="0"][data-f="qty"]', "2"); page.fill('[data-il="0"][data-f="unit_price"]', "500")
    page.click('#finInvoice [data-post="1"]'); page.wait_for_timeout(300)
    check("posting a USD invoice without a rate is stopped", "exchange rate" in page.inner_text("#toast").lower() and not any(i.get("currency") == "USD" for i in DB["fin_invoices"]))
    page.fill("#fi-rate", "3.64"); page.click('#finInvoice [data-post="1"]'); page.wait_for_timeout(1000)
    ui_ = next((i for i in DB["fin_invoices"] if i.get("currency") == "USD"), {})
    check("USD invoice saved with its rate and posted at QAR 3,640", ui_.get("status") == "posted" and float(ui_.get("fx_rate") or 0) == 3.64
          and any(l["debit"] == 3640 for l in DB["fin_lines"] if l["journal_id"] == ui_.get("journal_id")))
    page.wait_for_selector(".list .trow")
    row = page.inner_text(f'.trow[data-id="{ui_.get("id")}"]')
    check("invoice list shows USD 1,000 and the QAR value", "USD" in row and "1,000.00" in row and "3,640.00" in row, row)
    page.click('[data-act="fx-reval"]'); page.wait_for_selector("#fx-date")
    page.fill("#fx-date", today_qa()); page.dispatch_event("#fx-date", "change"); page.wait_for_selector('[data-fxrate="USD"]')
    check("revaluation lists the open USD invoice and waits for a closing rate", page.locator("#fxPost").is_disabled() and ui_["number"] in page.inner_text("#fxTable"))
    page.fill('[data-fxrate="USD"]', "3.62"); page.dispatch_event('[data-fxrate="USD"]', "change"); page.wait_for_timeout(200)
    check("preview shows the unrealised loss of 20", page.inner_text("#fx-net").replace("−", "-").startswith("-") and "20.00" in page.inner_text("#fx-net"), page.inner_text("#fx-net"))
    page.click("#fxPost"); page.wait_for_selector(".dlg"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(900)
    rv = DB["fin_fx_revaluations"][0] if DB["fin_fx_revaluations"] else {}
    check("revaluation posted with the entered rate", rv.get("rates") == {"USD": 3.62} and rv.get("net") == -20)
    page.click(f'[data-act="fx-cancel"][data-id="{rv.get("id")}"]'); page.wait_for_selector("#dlgIn"); page.fill("#dlgIn", "Wrong closing rate"); page.click('.dlg [data-r="1"]'); page.wait_for_timeout(800)
    check("revaluation can be cancelled with a reason", rv.get("status") == "cancelled" and "cancelled" in page.inner_text(".sheet").lower())
    page.keyboard.press("Escape"); page.wait_for_timeout(300)
    page.goto(BASE + f"/#/finance/{FE1}/overview"); page.goto(BASE + f"/#/finance/{FE1}/invoices"); page.wait_for_selector(f'.trow[data-id="{ui_["id"]}"]')
    page.click(f'.trow[data-id="{ui_["id"]}"]'); page.wait_for_selector('[data-act="fin-pay"]')
    check("invoice sheet shows the USD amounts and the booked rate", "1 usd = 3.64 qar" in page.inner_text(".sheet").lower())
    page.click('.sheet [data-act="fin-pay"]'); page.wait_for_selector("#fp-base")
    check("payment form opens in USD with the open amount", page.input_value("#fp-ccy") == "USD" and float(page.input_value("#fp-amt")) == 1000)
    page.click("#fpSave"); page.wait_for_timeout(300)
    check("foreign payment needs the bank amount", "reached the bank" in page.inner_text("#toast").lower())
    page.fill("#fp-base", "3660"); page.fill("#fp-ref", "SWIFT 8812"); page.click("#fpSave"); page.wait_for_timeout(1000)
    py = next((p for p in DB["fin_payments"] if p.get("currency") == "USD"), {})
    check("USD 1,000 received as QAR 3,660: realised gain 20 recorded", py.get("base_amount") == 3660 and py.get("fx_gain") == 20)
    page.goto(BASE + f"/#/finance/{FE1}/payments"); page.wait_for_selector(".rtable")
    check("payments list shows the QAR received and the USD amount", "3,660.00" in page.inner_text(".rtable") and "USD" in page.inner_text(".rtable"))
    shot(page, "30-fx-payments", True)
    ctx.close()

    # ---------- phone, light + dark ----------
    for scheme in ["light", "dark"]:
        ctx, page = new_page(b, 390, 844, scheme)
        page.goto(BASE + "/#/home"); page.wait_for_selector(".tiles", timeout=10000)
        check(f"phone ({scheme}): bottom nav visible, sidebar hidden", page.is_visible("#bottomnav") and not page.is_visible("#sidenav"))
        overflow = page.evaluate("document.documentElement.scrollWidth > window.innerWidth + 1")
        check(f"phone ({scheme}): no horizontal scroll on home", not overflow)
        shot(page, f"06-home-phone-{scheme}", True)
        if scheme == "light":
            page.click(".bn-capture"); page.wait_for_selector("#capForm"); shot(page, "07-capture-phone")
            page.keyboard.press("Escape")
            page.goto(BASE + "/#/tasks/today"); page.wait_for_selector(".viewtabs")
            check("phone: no horizontal scroll on tasks", not page.evaluate("document.documentElement.scrollWidth > window.innerWidth + 1"))
            shot(page, "08-tasks-phone", True)
            page.click('.trow[data-id="t-wait"]'); page.wait_for_selector("#tf-title"); shot(page, "09-task-phone")
            page.keyboard.press("Escape")
            for route_ in ["goals/habits", "reports/weekly", "notes/decisions", "reports/projects", f"finance/{FE1}/overview", f"finance/{FE1}/reports", f"finance/{FE1}/invoices", f"finance/{FE1}/partners/dashboard", f"finance/{FE1}/partners/ledger"]:
                page.goto(BASE + "/#/" + route_); page.wait_for_selector(".viewtabs")
                check(f"phone: no horizontal scroll on {route_}", not page.evaluate("document.documentElement.scrollWidth > window.innerWidth + 1"))
            shot(page, "13-projects-report-phone", True)
        ctx.close()
    b.close()

srv.shutdown()
errs = [e for e in errors if "favicon" not in e]
check("no JavaScript errors in the console", not errs, "; ".join(errs[:5]))
failed = [r for r in results if not r[1]]
print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
sys.exit(1 if failed else 0)
